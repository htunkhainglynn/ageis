from __future__ import annotations

"""Rebuild the long Aegis report using a chapter-first, two-appendix structure.

The factual source is the previously verified 148-page report.  Substantive
material is moved from appendices into Design, Implementation and Testing,
following the structural pattern of Hein Htet Aung's 2026 final report.
"""

from copy import deepcopy
from pathlib import Path
import importlib.util

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches


ROOT = Path("/Users/htunkhainglynn/Projects/aegis")
SOURCE = ROOT / "docs/report/Aegis_Project_Documentation_150_Pages.docx"
OUTPUT = ROOT / "docs/report/Aegis_Project_Documentation_Final_With_Verified_Evidence.docx"
OLD_BUILDER = ROOT / "docs/report/build_extended_project_documentation.py"
TIMELINE = ROOT / "docs/diagrams/aegis_dsdm_implementation_timeline.png"

spec = importlib.util.spec_from_file_location("old_report", OLD_BUILDER)
old = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(old)


def txt(element):
    return "".join(n.text or "" for n in element.iter(qn("w:t"))).strip()


def all_body_elements(doc):
    return [deepcopy(e) for e in doc._element.body if e.tag != qn("w:sectPr")]


def positions(elements):
    result = {}
    for i, element in enumerate(elements):
        value = txt(element)
        if value and value not in result:
            result[value] = i
    return result


def segment(elements, pos, start, end=None):
    a = pos[start]
    b = pos[end] if end else len(elements)
    return [deepcopy(e) for e in elements[a:b]]


def replace_text(element, value):
    nodes = list(element.iter(qn("w:t")))
    if not nodes:
        return
    nodes[0].text = value
    for node in nodes[1:]:
        node.text = ""


def set_style(element, style):
    ppr = element.find(qn("w:pPr"))
    if ppr is None:
        from docx.oxml import OxmlElement
        ppr = OxmlElement("w:pPr")
        element.insert(0, ppr)
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        from docx.oxml import OxmlElement
        pstyle = OxmlElement("w:pStyle")
        ppr.insert(0, pstyle)
    pstyle.set(qn("w:val"), style.replace(" ", ""))


def renumber(block, top_old, top_new, item_prefix_old, item_prefix_new):
    """Turn a former appendix into a chapter section and its items into H3."""
    for element in block:
        value = txt(element)
        if value == top_old:
            replace_text(element, top_new)
            set_style(element, "Heading 2")
        elif value.startswith(item_prefix_old):
            suffix = value[len(item_prefix_old):]
            replace_text(element, item_prefix_new + suffix)
            set_style(element, "Heading 3")
        elif f"Table {item_prefix_old}" in value or f"Extract {item_prefix_old}" in value:
            replace_text(element, value.replace(f"Table {item_prefix_old}", f"Table {item_prefix_new}")
                         .replace(f"Extract {item_prefix_old}", f"Extract {item_prefix_new}"))
        elif value.startswith("This appendix"):
            replace_text(element, "This section" + value[len("This appendix"):])
    return block


def rename_exact(block, old_value, new_value, style=None):
    for element in block:
        if txt(element) == old_value:
            replace_text(element, new_value)
            if style:
                set_style(element, style)
    return block


def renumber_chapter_block(block, old_number, new_number, new_title):
    """Renumber one complete chapter, including its headings and local captions."""
    for element in block:
        value = txt(element)
        updated = value
        if value.startswith(f"Chapter {old_number}:"):
            updated = new_title
            set_style(element, "Heading 1")
        elif value.startswith(f"{old_number}."):
            updated = f"{new_number}." + value[len(f"{old_number}."):]
        updated = updated.replace(f"Table {old_number}.", f"Table {new_number}.")
        updated = updated.replace(f"Figure {old_number}.", f"Figure {new_number}.")
        updated = updated.replace(f"Extract {old_number}.", f"Extract {new_number}.")
        if updated != value:
            replace_text(element, updated)
    return block


def remap_chapter_references(block):
    """Update prose cross-references without touching chapter-title paragraphs."""
    mapping = {4: 7, 5: 4, 6: 5, 7: 6, 8: 8, 9: 8, 10: 9}
    for element in block:
        # Rewrite one paragraph at a time.  A body element may be a table, in
        # which case operating on its concatenated text would destroy the cell
        # structure; traversing the cell paragraphs preserves it.
        paragraphs = [element] if element.tag == qn("w:p") else list(element.iter(qn("w:p")))
        for paragraph in paragraphs:
            value = txt(paragraph)
            if not value or value.startswith("Chapter "):
                continue
            updated = value
            for old_number in mapping:
                updated = updated.replace(f"Chapter {old_number}", f"Chapter @@{old_number}@@")
            for old_number, new_number in mapping.items():
                updated = updated.replace(f"Chapter @@{old_number}@@", f"Chapter {new_number}")
            if updated != value:
                replace_text(paragraph, updated)
    return block


def integrate_testing_chapter(block):
    """Move former Chapter 9 content under Chapter 8 as section 8.18."""
    for element in block:
        value = txt(element)
        updated = value
        if value == "Chapter 9: Testing":
            updated = "8.18 Integrated Verification Evidence"
            set_style(element, "Heading 2")
        elif value.startswith("9."):
            updated = "8.18." + value[2:]
            set_style(element, "Heading 3")
        updated = updated.replace("Table 9.1", "Table 8.8").replace("Table 9.2", "Table 8.9")
        updated = (updated.replace("Figure 9.1", "Figure 8.23")
                   .replace("Figure 9.2", "Figure 8.24")
                   .replace("Figure 9.3", "Figure 8.25"))
        if updated != value:
            replace_text(element, updated)
    return block


def create_timeline():
    from datetime import datetime
    from PIL import Image, ImageDraw, ImageFont

    rows = [
        ("Design (not timeboxed)", "2026-04-27", "2026-05-29", "#D9D9D9"),
        ("TB1 Foundation and identity", "2026-06-01", "2026-06-17", "#222222"),
        ("TB2 Keys and policy", "2026-06-18", "2026-07-03", "#3D3D3D"),
        ("TB3 Proxy authentication", "2026-07-06", "2026-07-22", "#555555"),
        ("TB4 Limits and threats", "2026-07-23", "2026-08-07", "#6E6E6E"),
        ("TB5 Distribution and dashboard", "2026-08-10", "2026-08-26", "#858585"),
        ("TB6 Integration and hardening", "2026-08-27", "2026-09-11", "#999999"),
        ("Testing (not timeboxed)", "2026-09-14", "2026-10-09", "#D9D9D9"),
        ("Evaluation (not timeboxed)", "2026-10-12", "2026-10-30", "#D9D9D9"),
    ]
    width, height = 2400, 1040
    label_left, left, right, top, bottom = 30, 720, 45, 150, 135
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    font = ImageFont.truetype(font_path, 27)
    small = ImageFont.truetype(font_path, 24)
    bold = ImageFont.truetype(bold_path, 36)
    draw.text((width // 2, 35), "Aegis project schedule: DSDM timeboxes apply only to implementation",
              font=bold, fill="black", anchor="ma")
    start_all = datetime.fromisoformat("2026-04-27")
    end_all = datetime.fromisoformat("2026-10-30")
    days = (end_all - start_all).days
    plot_w = width - left - right
    plot_h = height - top - bottom
    row_h = plot_h / len(rows)

    def x_for(day):
        return left + int(((day - start_all).days / days) * plot_w)

    tick = start_all
    while tick <= end_all:
        x = x_for(tick)
        draw.line((x, top - 10, x, height - bottom + 8), fill="#D0D0D0", width=2)
        draw.text((x, height - bottom + 25), tick.strftime("%d %b"), font=small, fill="black", anchor="ma")
        from datetime import timedelta
        tick += timedelta(days=14)

    for y_index, (label, start, end, colour) in enumerate(rows):
        y0 = top + int(y_index * row_h + row_h * .18)
        y1 = top + int((y_index + 1) * row_h - row_h * .18)
        x0 = x_for(datetime.fromisoformat(start))
        x1 = x_for(datetime.fromisoformat(end))
        draw.rectangle((x0, y0, x1, y1), fill=colour, outline="black", width=2)
        draw.text((label_left, (y0 + y1) // 2), label, font=font, fill="black", anchor="lm")
    draw.line((left, height - bottom + 8, width - right, height - bottom + 8), fill="black", width=2)
    draw.text((width // 2, height - 18), "2026", font=font, fill="black", anchor="ms")
    TIMELINE.parent.mkdir(parents=True, exist_ok=True)
    image.save(TIMELINE, dpi=(220, 220))


def toc_page(doc, title, entries):
    doc.add_paragraph(title, style="Front Matter Heading")
    for label, key, level in entries:
        old.toc_entry(doc, label, f"@@{key}@@", level)
    old.page_break(doc)


def front_matter(doc):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    old.set_run(p.add_run("COMP 1682 PROJECT DOCUMENTATION"), size=16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = old.Pt(16)
    old.set_run(p.add_run("Aegis"), size=16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    old.set_run(p.add_run("API Security and Management System"), size=14, bold=True)
    for _ in range(2):
        doc.add_paragraph()
    for value, size, bold in [
        ("Htun Khaing Lynn", 12, True),
        ("Student ID: 001557481", 11, False),
        ("BSc (Hons) Computing", 11, False),
        ("Final Project Report - 2026", 11, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        old.set_run(p.add_run(value), size=size, bold=bold)
    old.page_break(doc)

    doc.add_paragraph("Abstract", style="Front Matter Heading")
    old.para(doc, "Aegis is a self-hosted API security and management system that centralises authentication, authorisation, rate limiting, threat detection and audit visibility in front of protected services. It combines a Python/FastAPI Control Plane, a Go Reverse Proxy and a React dashboard with PostgreSQL and Redis. This report evaluates the implemented system against the Project Proposal, current source code, database migrations and verified local test evidence.")
    old.para(doc, "The implementation issues cryptographically random API keys whose hashes alone are persisted; validates API keys and optional JWTs; binds exact HTTP method/path pairs to required resource:action scopes; applies fixed-window, sliding-window and token-bucket limits in Redis; distributes policy by REST bootstrap and authenticated gRPC streaming; detects configured request-target patterns; enforces manual and automatic IP blocks; and records sanitised analytics. The report explains how JSON payload structures are parsed and validated and how the Go proxy consumes the resulting policy without inferring semantics from scope names.")
    old.para(doc, "Verification includes 47 Control Plane tests at 77% application coverage, Go race/vet/format checks and 83.4% enforcement-core coverage, six dashboard tests, thirteen POC/tooling tests, focused request-schema and exact-scope tests, an observed live attack/defence flow and a completed isolated Docker E2E run. In line with the structural reference, detailed design, implementation, operation and testing material is presented inside the relevant chapters. Only the proposal pointer and project-plan figure remain as appendices.")
    old.para(doc, "Keywords: API security, reverse proxy, FastAPI, Go, JWT, rate limiting, Redis, DSDM.", italic=True)
    old.page_break(doc)

    doc.add_paragraph("Acknowledgements", style="Front Matter Heading")
    old.para(doc, "I would like to thank the project supervisor, module tutors and COMP 1682 teaching team for their guidance on scope, academic writing, evidence and project management. Their feedback led to a clearer distinction between implementation delivery timeboxes and the surrounding design, testing and evaluation activities.")
    old.para(doc, "The report also acknowledges the standards, research literature and open-source communities referenced in the bibliography. Hein Htet Aung's 2026 final report was used as the latest structural reference: its chapter-first treatment of implementation, testing and deployment, and its restrained appendix use, informed this revision. Its project-specific content was not reused as Aegis evidence.")
    old.page_break(doc)

    toc_page(doc, "Table of Contents", [
        ("Abstract", "ABSTRACT", 0), ("Acknowledgements", "ACK", 0),
        ("List of Figures", "FIGURES", 0), ("List of Tables", "TABLES", 0),
        ("Chapter 1: Introduction", "CH1", 0), ("1.1-1.9 Background, problem, aim, objectives and scope", "CH1", 1),
        ("Chapter 2: Product Research", "CH2", 0), ("2.1-2.6 Product comparison and research findings", "CH2", 1),
        ("Chapter 3: Literature Review", "CH3", 0), ("3.1-3.7 Security, identity, rate limiting and architecture", "CH3", 1),
        ("Chapter 4: Legal, Social, Ethical and Professional Issues", "CH4", 0),
        ("4.1-4.7 Legal, social, ethical and professional analysis", "CH4", 1),
    ])
    toc_page(doc, "Table of Contents (continued)", [
        ("Chapter 5: System Analysis", "CH5", 0), ("5.1-5.6 Users, requirements and MoSCoW priorities", "CH5", 1),
        ("Chapter 6: System Design", "CH6", 0), ("6.1-6.9 Architecture, domain model and enforcement design", "CH6", 1),
        ("6.10 Detailed domain data dictionary", "DESIGN_DATA", 1),
        ("6.11 Reverse proxy enforcement walkthrough", "DESIGN_PROXY", 1),
        ("Chapter 7: Project Timeline", "CH7", 0), ("7.1-7.12 DSDM rules and six implementation-only timeboxes", "CH7", 1),
        ("Chapter 8: System Implementation", "CH8", 0), ("8.1-8.11 Design, coding and testing evidence for six timeboxes", "CH8", 1),
        ("8.12 DSDM review and sign-off records", "IMPL_DSDM", 1), ("8.13 Valid JSON payloads", "IMPL_JSON", 1),
        ("8.14 Control Plane API operation catalogue", "IMPL_API", 1),
        ("8.15 Operator and API consumer procedures", "IMPL_GUIDE", 1),
        ("8.16 Deployment and operations", "IMPL_OPS", 1), ("8.17 Reverse-learning Go and JSON", "IMPL_LEARN", 1),
    ])
    toc_page(doc, "Table of Contents (continued)", [
        ("8.18 Integrated verification evidence", "TEST_CORE", 1),
        ("8.19 Selected test case matrix", "TEST_MATRIX", 1), ("8.20 Evidence and reproducibility", "TEST_EVIDENCE", 1),
        ("8.21 Detailed testing and traceability", "TEST_DETAIL", 1),
        ("Chapter 9: Evaluation and Reflection", "CH9", 0),
        ("9.1-9.8 Objectives, limitations, reflection and conclusion", "CH9", 1),
        ("References", "REFERENCES", 0), ("Appendices", "APPENDICES", 0),
        ("Appendix A: Project Proposal", "APPENDICES", 1), ("Appendix B: Figures", "APPENDICES", 1),
    ])

    doc.add_paragraph("List of Figures", style="Front Matter Heading")
    old.table(doc, ["Figure", "Description"], [
        ("Figure 5.1", "Aegis use-case model"), ("Figure 6.1", "Mermaid current-system architecture"),
        ("Figure 6.2", "Mermaid persisted domain model"), ("Figure 6.3", "Mermaid Control Plane validation and persistence sequence"),
        ("Figure 6.4", "Mermaid Reverse Proxy enforcement sequence"), ("Figure 6.5", "Mermaid policy distribution and resilience flow"),
        ("Figure 6.6", "Mermaid rate-limit decision flow"), ("Figure 6.7", "Mermaid security-event and analytics pipeline"),
        ("Figure 8.1", "TB1 identity and RBAC test evidence"), ("Figure 8.2", "Successful typed-login and isolated-stack evidence"),
        ("Figure 8.3", "Authentication and RBAC rejection evidence"), ("Figure 8.4", "Actual dashboard login interface"),
        ("Figure 8.5", "TB2 key and policy test evidence"), ("Figure 8.6", "API-key creation and one-time secret evidence"),
        ("Figure 8.7", "Invalid route-permission policy validation evidence"), ("Figure 8.8", "Generated FastAPI OpenAPI interface"),
        ("Figure 8.9", "TB3 proxy authentication and authorisation tests"), ("Figure 8.10", "Exact-scope allow evidence"),
        ("Figure 8.11", "Insufficient-scope denial evidence"), ("Figure 8.12", "TB4 limits, threat and IP-enforcement tests"),
        ("Figure 8.13", "Redis rate-limit sequence evidence"), ("Figure 8.14", "Threat and exact-IP enforcement evidence"),
        ("Figure 8.15", "TB5 distribution, analytics and dashboard tests"), ("Figure 8.16", "Authenticated policy-distribution evidence"),
        ("Figure 8.17", "Security analytics and automatic-block evidence"), ("Figure 8.18", "TB6 complete quality-gate evidence"),
        ("Figure 8.19", "Integrated service health evidence"), ("Figure 8.20", "Completed end-to-end allow and deny evidence"),
        ("Figure 8.21", "Actual attack/defence console"), ("Figure 8.22", "Observed live POC output"),
        ("Figure 8.23", "Completed automated quality-gate result"), ("Figure 8.24", "Completed live POC enforcement result"),
        ("Figure 8.25", "Completed isolated Docker end-to-end result"),
        ("Figure B.1", "Project plan showing implementation-only DSDM timeboxes"),
    ], [1.2, 5.3])
    old.page_break(doc)

    doc.add_paragraph("List of Tables", style="Front Matter Heading")
    old.table(doc, ["Chapter", "Principal table groups"], [
        ("1-3", "SWOT, product comparison, algorithms and literature synthesis."),
        ("4-6", "LSEP mitigations, requirements, components, data definitions and enforcement design."),
        ("7", "DSDM rules, MoSCoW balance and dated From/To/Duration/Summary schedules."),
        ("8", "Timebox design, coding and testing evidence; payloads; operations; deployment and traceability."),
        ("9", "Objective evaluation, limitations, reflection and conclusion."),
    ], [1.1, 5.4])


def append_new(doc, elements):
    old.append_elements(doc, elements)


def add_bridge(doc, title, body):
    old.heading(doc, title, 2)
    old.para(doc, body)


def add_appendices(doc):
    old.heading(doc, "Appendices", 1)
    old.heading(doc, "Appendix A: Project Proposal", 2)
    old.para(doc, "The approved Aegis Project Proposal is submitted as the accompanying source document. It defines the original problem, aim, objectives, scope, risk register and overall schedule. Chapters 1, 5, 7, 8 and 9 explicitly trace the implemented system back to that baseline and identify material adaptations rather than presenting them as if they were planned from the outset.")
    old.table(doc, ["Proposal baseline", "How this report uses it"], [
        ("Design: 27 April-29 May 2026", "Treated as a project activity, not as a DSDM delivery timebox."),
        ("Implementation: 1 June-11 September 2026", "Divided into six fixed DSDM structured timeboxes."),
        ("Testing: 14 September-9 October 2026", "Retained as a separate scheduled activity, with its evidence integrated into Chapter 8 beside the relevant implementation timeboxes."),
        ("Evaluation: 12-30 October 2026", "Used to assess delivery, limitations and learning against the original objectives."),
    ], [2.2, 4.3], "Table A.1: Proposal schedule interpretation used throughout the report.")
    old.heading(doc, "Appendix B: Figures", 2)
    old.para(doc, "Figure B.1 reproduces the consolidated project plan. Dark bars are the six implementation delivery timeboxes. Design, formal testing and evaluation remain visible as project activities but are deliberately labelled as not timeboxed, preserving the DSDM boundary required for this report.")
    old.add_figure(doc, TIMELINE, "Figure B.1: Aegis project plan with DSDM timeboxes applied only to implementation.", width=6.25)


def materialise_toc_pages(doc):
    page_map = {
        "ABSTRACT": 2, "ACK": 3, "FIGURES": 7, "TABLES": 9,
        "CH1": 10, "CH2": 15, "CH3": 18, "CH4": 20, "CH5": 23, "CH6": 32,
        "DESIGN_DATA": 46, "DESIGN_PROXY": 54, "CH7": 65,
        "CH8": 80, "IMPL_DSDM": 96, "IMPL_JSON": 101, "IMPL_API": 103,
        "IMPL_GUIDE": 118, "IMPL_OPS": 126, "IMPL_LEARN": 136,
        "TEST_CORE": 143, "TEST_MATRIX": 146, "TEST_EVIDENCE": 147, "TEST_DETAIL": 148,
        "CH9": 160, "REFERENCES": 163, "APPENDICES": 165,
    }
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for key, page in page_map.items():
                marker = f"@@{key}@@"
                if marker in run.text:
                    run.text = run.text.replace(marker, str(page))
                    old.set_run(run, size=11, bold=run.bold)


def build():
    create_timeline()
    doc = Document(SOURCE)
    elements = all_body_elements(doc)
    pos = positions(elements)

    front_end = pos["Chapter 1: Introduction"]
    ch1_to_3 = segment(elements, pos, "Chapter 1: Introduction", "Chapter 4: Project Timeline")
    timeline = segment(elements, pos, "Chapter 4: Project Timeline", "Chapter 5: Legal, Social, Ethical and Professional (LSEP) Issues")
    lsep = segment(elements, pos, "Chapter 5: Legal, Social, Ethical and Professional (LSEP) Issues", "Chapter 6: System Analysis")
    analysis = segment(elements, pos, "Chapter 6: System Analysis", "Chapter 7: System Design")
    design = segment(elements, pos, "Chapter 7: System Design", "Chapter 8: System Implementation")
    ch8 = segment(elements, pos, "Chapter 8: System Implementation", "Chapter 9: Testing")
    testing = segment(elements, pos, "Chapter 9: Testing", "Chapter 10: Evaluation and Reflection")
    evaluation = segment(elements, pos, "Chapter 10: Evaluation and Reflection", "References")
    refs = segment(elements, pos, "References", "Appendix A: Valid Example JSON Payloads")

    for block in (ch1_to_3, timeline, lsep, analysis, design, ch8, testing, evaluation):
        remap_chapter_references(block)
    timeline = renumber_chapter_block(timeline, 4, 7, "Chapter 7: Project Timeline")
    lsep = renumber_chapter_block(lsep, 5, 4, "Chapter 4: Legal, Social, Ethical and Professional (LSEP) Issues")
    analysis = renumber_chapter_block(analysis, 6, 5, "Chapter 5: System Analysis")
    design = renumber_chapter_block(design, 7, 6, "Chapter 6: System Design")
    evaluation = renumber_chapter_block(evaluation, 10, 9, "Chapter 9: Evaluation and Reflection")
    testing = integrate_testing_chapter(testing)

    json_payloads = renumber(segment(elements, pos, "Appendix A: Valid Example JSON Payloads", "Appendix B: Selected Test Case Matrix"),
                             "Appendix A: Valid Example JSON Payloads", "8.13 Valid Example JSON Payloads", "A.", "8.13.")
    test_matrix = renumber(segment(elements, pos, "Appendix B: Selected Test Case Matrix", "Appendix C: Evidence and Reproducibility Index"),
                           "Appendix B: Selected Test Case Matrix", "8.19 Selected Test Case Matrix", "B.", "8.19.")
    evidence = renumber(segment(elements, pos, "Appendix C: Evidence and Reproducibility Index", "Appendix D: DSDM Implementation Timebox Review Records"),
                        "Appendix C: Evidence and Reproducibility Index", "8.20 Evidence and Reproducibility", "C.", "8.20.")
    dsdm = renumber(segment(elements, pos, "Appendix D: DSDM Implementation Timebox Review Records", "Appendix E: Control Plane API Operation Catalogue"),
                     "Appendix D: DSDM Implementation Timebox Review Records", "8.12 DSDM Implementation Timebox Review Records", "D.", "8.12.")
    for element in dsdm:
        value = txt(element)
        if value.startswith("These records provide the implementation depth used by the reference report"):
            replace_text(element, value.replace(
                "These records provide the implementation depth used by the reference report",
                "These records provide detailed retrospective implementation evidence",
            ))
    api = renumber(segment(elements, pos, "Appendix E: Control Plane API Operation Catalogue", "Appendix F: Domain Data Dictionary"),
                    "Appendix E: Control Plane API Operation Catalogue", "8.14 Control Plane API Operation Catalogue", "E.", "8.14.")
    data = renumber(segment(elements, pos, "Appendix F: Domain Data Dictionary", "Appendix G: Operator and API Consumer User Manual"),
                     "Appendix F: Domain Data Dictionary", "6.10 Detailed Domain Data Dictionary", "F.", "6.10.")
    guide = renumber(segment(elements, pos, "Appendix G: Operator and API Consumer User Manual", "Appendix H: Reverse Proxy Enforcement Walkthrough"),
                      "Appendix G: Operator and API Consumer User Manual", "8.15 Operator and API Consumer Procedures", "G.", "8.15.")
    proxy = renumber(segment(elements, pos, "Appendix H: Reverse Proxy Enforcement Walkthrough", "Appendix I: Detailed Testing and Traceability Records"),
                      "Appendix H: Reverse Proxy Enforcement Walkthrough", "6.11 Reverse Proxy Enforcement Walkthrough", "H.", "6.11.")
    tests = renumber(segment(elements, pos, "Appendix I: Detailed Testing and Traceability Records", "Appendix J: Deployment and Operations Runbook"),
                      "Appendix I: Detailed Testing and Traceability Records", "8.21 Detailed Testing and Traceability Records", "I.", "8.21.")
    ops = renumber(segment(elements, pos, "Appendix J: Deployment and Operations Runbook", "Appendix K: Reverse-Learning Notes for Go and JSON"),
                    "Appendix J: Deployment and Operations Runbook", "8.16 Deployment and Operations", "J.", "8.16.")
    learn = renumber(segment(elements, pos, "Appendix K: Reverse-Learning Notes for Go and JSON", "Supplementary References"),
                      "Appendix K: Reverse-Learning Notes for Go and JSON", "8.17 Reverse-Learning Notes for Go and JSON", "K.", "8.17.")
    supp = segment(elements, pos, "Supplementary References")
    # Merge the supplementary source entries into References without a second top-level heading.
    supp = [e for e in supp if txt(e) != "Supplementary References"]

    rename_exact(design, "6.9 Chapter Summary", "6.9 Core Design Summary")
    rename_exact(ch8, "8.11 Chapter Summary", "8.11 Core Implementation Summary")
    rename_exact(testing, "8.18.8 Chapter Summary", "8.18.8 Core Verification Summary")

    old.clear_body(doc)
    old.configure_styles(doc)
    front_matter(doc)
    append_new(doc, ch1_to_3)
    append_new(doc, lsep)
    append_new(doc, analysis)
    append_new(doc, design)
    append_new(doc, data)
    append_new(doc, proxy)
    add_bridge(doc, "6.12 Chapter Summary", "The design chapter contains the persisted data definitions and the complete enforcement walkthrough at the point where they are needed. Project Timeline follows this architecture chapter in direct response to tutor feedback.")
    append_new(doc, timeline)
    append_new(doc, ch8)
    append_new(doc, dsdm)
    append_new(doc, json_payloads)
    append_new(doc, api)
    append_new(doc, guide)
    append_new(doc, ops)
    append_new(doc, learn)
    append_new(doc, testing)
    append_new(doc, test_matrix)
    append_new(doc, evidence)
    append_new(doc, tests)
    add_bridge(doc, "8.22 Chapter Summary", "System Implementation now combines timebox design decisions, representative code, focused and complete test runs, screenshots, live POC outcomes, traceability and acknowledged gaps. No failed or incomplete environment run is represented as a pass.")
    append_new(doc, evaluation)
    append_new(doc, refs)
    append_new(doc, supp)
    old.para(doc, "Hein Htet Aung (2026) An AI-Driven Retrieval-Augmented Generation (RAG) System for Enhancing Entrepreneurial Knowledge Acquisition among Young Adults: Final Report. Unpublished COMP 1682 final project report.", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_appendices(doc)

    old.configure_pages_and_furniture(doc)
    old.normalise_direct_formatting(doc)
    materialise_toc_pages(doc)
    doc.core_properties.title = "Aegis API Security and Management System - Reference-Aligned Final Report"
    doc.core_properties.subject = "COMP 1682 chapter-first project documentation"
    doc.core_properties.author = "Htun Khaing Lynn"
    doc.core_properties.keywords = "Aegis, API security, reverse proxy, DSDM, final report"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
