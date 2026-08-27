from __future__ import annotations

"""Create the compact Aegis project report requested for final submission.

The verified long-form report remains the factual baseline.  This final-stage
builder replaces Chapters 6-8 with diagram-led material, presents each Chapter
7 timebox as a simple date/task/deliverable record, and formats every
implementation timebox in Chapter 8 as Design, Coding, Testing and Retro.
"""

from pathlib import Path
import importlib.util
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path("/Users/htunkhainglynn/Projects/aegis")
SOURCE = ROOT / "docs/report/Aegis_Project_Documentation_Final_With_Verified_Evidence.docx"
OUTPUT = ROOT / "docs/report/Aegis_Project_Documentation_Compact.docx"
DIAGRAMS = ROOT / "docs/diagrams"
IMAGES = ROOT / "docs/evidence"
HELPERS = ROOT / "docs/report/build_extended_project_documentation.py"

spec = importlib.util.spec_from_file_location("report_helpers", HELPERS)
helpers = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(helpers)


TOC_PAGES = {
    "Abstract": "2",
    "Acknowledgement": "3",
    "List of Figures": "5",
    "List of Tables": "6",
    "Chapter 1: Introduction": "7",
    "Chapter 2: Product Research": "12",
    "Chapter 3: Literature Review": "15",
    "Chapter 4: Legal, Social, Ethical and Professional Issues": "17",
    "Chapter 5: System Analysis": "20",
    "Chapter 6: System Design": "29",
    "Chapter 7: Project Timeline": "35",
    "Chapter 8: System Implementation": "45",
    "Chapter 9: Evaluation and Reflection": "62",
    "References": "65",
    "Appendices": "67",
}


TIMEBOXES = [
    {
        "id": "TB1", "section": "8.2", "name": "Control Plane Foundation and Identity",
        "dates": "1-17 June 2026",
        "design": "The first increment established the trusted administrative boundary. The User model stores a UUID, unique email, password hash, role, active state and timestamps. FastAPI request models validate typed input before a service applies registration and role rules. PostgreSQL is the durable source for identity, while Redis stores refresh-token revocation state. Public registration is restricted to API Consumer accounts; Admin and Viewer roles are assigned through protected administration routes.",
        "coding": "The implementation separates routers, Pydantic schemas, services and repositories. Passwords are hashed before persistence; login returns short-lived access and refresh tokens; refresh checks revocation; logout records the refresh-token identifier. Dependencies resolve the current user and required role before protected operations run.",
        "testing": "Focused authentication and RBAC tests covered registration, login, refresh, logout, bootstrap administration and Admin-only user management. The current quality run completed 47 Control Plane tests. Invalid credentials return a controlled 401 response, an unauthorised role returns 403, and public registration cannot create an Admin account.",
        "retro": "Separating input validation from business rules made failures easier to locate. The main lesson was that role values in JSON are untrusted input: the server, not the browser, must decide which roles a caller may create or assign.",
    },
    {
        "id": "TB2", "section": "8.3", "name": "API Keys and Administrative Policy",
        "dates": "18 June-3 July 2026",
        "design": "An API key is represented by durable metadata plus a one-way hash. The raw secret is returned only after creation. Scope labels use resource:action strings, for example echo:read, but the label has no routing meaning by itself. A RoutePermission binds an HTTP method and path pattern to the required label. Administrative schemas also validate rate limits, JWT configurations, threat rules and exact IP blocks.",
        "coding": "The service generates a cryptographically random key, stores its bcrypt hash and prefix, and returns the raw value once. List responses expose only masked metadata. Revocation soft-disables a key so historical events remain traceable. Policy services reject invalid algorithms, duplicate active route scopes and references to missing keys.",
        "testing": "Scenario tests confirmed one-time secret disclosure, owner and Admin access rules, soft revocation, duplicate-policy rejection and typed validation errors. Invalid combinations are rejected before database writes, while a revoked key remains identifiable but cannot authenticate proxy traffic.",
        "retro": "The initial idea that a scope alone identifies a backend was too ambiguous. The corrected design treats the scope as a permission label and makes the route-permission record the explicit connection between a real request and that label.",
    },
    {
        "id": "TB3", "section": "8.4", "name": "Reverse Proxy Authentication and Authorisation",
        "dates": "6-22 July 2026",
        "design": "The Go handler enforces policy before the standard reverse proxy forwards a request. The ordered path is request context, direct-peer IP block, threat matching, API-key validation, optional JWT validation, exact route permission, Redis rate limit, credential stripping and upstream forwarding. Failures stop immediately and use stable JSON error codes.",
        "coding": "Go interfaces isolate the Control Plane client, policy provider, limit store and event reporter, so handler branches can be tested without a full stack. The key cache stores successful and negative lookups for bounded periods. JWT validation pins the configured algorithm and checks expiry. Before forwarding, Aegis API-key and authorization headers are removed.",
        "testing": "Go tests verified missing, malformed, revoked and invalid keys; JWT expiry and algorithm substitution; exact-scope matching; denied write access; header stripping; and successful upstream forwarding. The focused scope run showed echo:read allows the configured GET route, whereas a missing echo:write denies POST and never calls the backend.",
        "retro": "Middleware order is part of the security design. Cheap deterministic rejections should occur early, but each denial must still produce a sanitised event. Exact matching is intentionally conservative; wildcard-looking scope labels do not grant implicit permissions.",
    },
    {
        "id": "TB4", "section": "8.5", "name": "Distributed Limits and Threat Enforcement",
        "dates": "23 July-7 August 2026",
        "design": "Redis provides shared counters for fixed-window, sliding-window and token-bucket rules. Precedence chooses the most specific active rule: API-key, then route, then global. Exact direct-peer IP blocks and RE2-compatible request-target patterns run before credential validation. Repeated qualifying events can create an automatic block through the Control Plane.",
        "coding": "Lua scripts perform atomic counter updates and return the decision plus reset information. The proxy emits 429 with a stable RATE_LIMIT_EXCEEDED code when capacity is exhausted. Threat patterns are compiled safely, exact IP values are canonicalised, and automatic blocks are persisted rather than remaining only in one proxy process.",
        "testing": "Unit and integration tests covered counter sharing, key isolation, refill and expiry, invalid threat expressions, direct-peer IP handling and race safety. The isolated stack produced 200, 200, 429 for a two-request fixed window, returned 403 for a matching threat, returned 403 for an exact IP block, and confirmed that an automatic block persisted in Control Plane data.",
        "retro": "A stale loopback block affected an early local run. The proxy was behaving correctly; the environment was not clean. The timebox therefore added explicit setup checks and disposable data to distinguish policy state from implementation defects.",
    },
    {
        "id": "TB5", "section": "8.6", "name": "Policy Distribution, Analytics and Dashboard",
        "dates": "10-26 August 2026",
        "design": "The proxy obtains a REST bootstrap snapshot, then subscribes to authenticated gRPC updates. A last-valid in-memory copy continues to protect traffic during a temporary Control Plane interruption. A bounded asynchronous reporter sends sanitised security events. The React dashboard is role-aware and keeps its access token only in memory.",
        "coding": "The subscriber authenticates with an internal token, swaps complete snapshots atomically and reconnects with backoff. The reporter excludes raw API keys, JWTs and bodies. Control Plane analytics aggregate forwarded and blocked outcomes. Dashboard routes cover login, API keys, policies, IP blocks, threat rules and analytics according to the current role.",
        "testing": "Tests covered rejected internal tokens, bootstrap and stream updates, stale-policy fallback, asynchronous flush, dashboard route guards and live analytics data. The isolated run activated an HS256 policy, delivered it through authenticated gRPC, forwarded a 200 request, then showed forwarded, blocked, threat and rate-limit events in analytics.",
        "retro": "The key learning was to move whole immutable snapshots rather than mutate shared policy structures field by field. For the browser, memory-only tokens reduce persistence risk but require a new login after a reload, which is an accepted usability trade-off.",
    },
    {
        "id": "TB6", "section": "8.7", "name": "Integration, POC and Implementation Hardening",
        "dates": "27 August-11 September 2026",
        "design": "The final implementation increment connected PostgreSQL, Redis, FastAPI, gRPC, the Go proxy, the dashboard and a local echo upstream through Docker Compose. Generated local credentials, health checks and deterministic setup scripts support repeatable demonstrations without storing secrets in the report.",
        "coding": "Compose wiring defines component dependencies and local ports. The end-to-end script creates disposable data, logs in, creates policies and keys, exercises allow and deny paths, checks analytics and cleans up the isolated stack. The attack/defence console offers the same core scenarios through a safe localhost-only web interface.",
        "testing": "The complete quality gate passed: 47 Control Plane tests, Go format/vet/race checks with 83.4% proxy-core coverage, six dashboard tests and thirteen POC/tooling tests. The isolated Docker run completed health, authenticated policy delivery, upstream forwarding, invalid and revoked credential rejection, manual and automatic IP blocking, threat matching, rate limiting and analytics.",
        "retro": "Cross-component tests found configuration and state problems that unit tests could not reveal. The final process now uses generated credentials, redacted output, disposable resources and explicit cleanup. Production load testing and an independent penetration review remain future work.",
    },
]


TIMEBOX_SIMPLE_SCHEDULES = {
    "TB1": [
        ("Implementation Setup", "0.5 day", "1 Jun 2026", "1 Jun 2026"),
        ("Feature Build", "2 days", "1 Jun 2026", "3 Jun 2026"),
        ("Testing and Fixing", "8 days", "3 Jun 2026", "15 Jun 2026"),
        ("Integration and Documentation", "2 days", "15 Jun 2026", "17 Jun 2026"),
        ("Review and Sign-off", "0.5 day", "17 Jun 2026", "17 Jun 2026"),
    ],
    "TB2": [
        ("Implementation Setup", "0.5 day", "18 Jun 2026", "18 Jun 2026"),
        ("Feature Build", "2 days", "18 Jun 2026", "22 Jun 2026"),
        ("Testing and Fixing", "7 days", "22 Jun 2026", "1 Jul 2026"),
        ("Integration and Documentation", "2 days", "1 Jul 2026", "3 Jul 2026"),
        ("Review and Sign-off", "0.5 day", "3 Jul 2026", "3 Jul 2026"),
    ],
    "TB3": [
        ("Implementation Setup", "0.5 day", "6 Jul 2026", "6 Jul 2026"),
        ("Feature Build", "2 days", "6 Jul 2026", "8 Jul 2026"),
        ("Testing and Fixing", "8 days", "8 Jul 2026", "20 Jul 2026"),
        ("Integration and Documentation", "2 days", "20 Jul 2026", "22 Jul 2026"),
        ("Review and Sign-off", "0.5 day", "22 Jul 2026", "22 Jul 2026"),
    ],
    "TB4": [
        ("Implementation Setup", "0.5 day", "23 Jul 2026", "23 Jul 2026"),
        ("Feature Build", "2 days", "23 Jul 2026", "27 Jul 2026"),
        ("Testing and Fixing", "7 days", "27 Jul 2026", "5 Aug 2026"),
        ("Integration and Documentation", "2 days", "5 Aug 2026", "7 Aug 2026"),
        ("Review and Sign-off", "0.5 day", "7 Aug 2026", "7 Aug 2026"),
    ],
    "TB5": [
        ("Implementation Setup", "0.5 day", "10 Aug 2026", "10 Aug 2026"),
        ("Feature Build", "2 days", "10 Aug 2026", "12 Aug 2026"),
        ("Testing and Fixing", "8 days", "12 Aug 2026", "24 Aug 2026"),
        ("Integration and Documentation", "2 days", "24 Aug 2026", "26 Aug 2026"),
        ("Review and Sign-off", "0.5 day", "26 Aug 2026", "26 Aug 2026"),
    ],
    "TB6": [
        ("Implementation Setup", "0.5 day", "27 Aug 2026", "27 Aug 2026"),
        ("Feature Build", "2 days", "27 Aug 2026", "31 Aug 2026"),
        ("Testing and Fixing", "7 days", "31 Aug 2026", "9 Sep 2026"),
        ("Integration and Documentation", "2 days", "9 Sep 2026", "11 Sep 2026"),
        ("Review and Sign-off", "0.5 day", "11 Sep 2026", "11 Sep 2026"),
    ],
}


TIMEBOX_SIMPLE_DELIVERABLES = {
    "TB1": {
        "Design": [
            "User, role and session data design.",
            "Authentication flow and database migration plan.",
        ],
        "Coding": [
            "FastAPI registration, login, refresh and logout.",
            "Persisted RBAC, Admin user management and Redis token revocation.",
        ],
        "Non-Functional Requirements": [
            "Hashed passwords, typed request validation and server-side role enforcement.",
        ],
        "Testing": [
            "Migration, authentication, token lifecycle and RBAC tests.",
        ],
    },
    "TB2": {
        "Design": [
            "API-key lifecycle and administrative policy structure.",
            "Exact method/path to resource:action scope mapping.",
        ],
        "Coding": [
            "One-time key issuance, bcrypt persistence, prefix lookup and soft revocation.",
            "CRUD for rate limits, JWT settings, threat rules, IP blocks and route permissions.",
        ],
        "Non-Functional Requirements": [
            "Secret non-disclosure, ownership checks and duplicate-active-policy protection.",
        ],
        "Testing": [
            "Key ownership, revocation, validation and policy-conflict tests.",
        ],
    },
    "TB3": {
        "Design": [
            "Ordered proxy enforcement flow and Control Plane validation contract.",
        ],
        "Coding": [
            "Go reverse proxy, API-key cache, JWT validation and exact scope authorisation.",
            "Credential stripping and structured proxy error responses.",
        ],
        "Non-Functional Requirements": [
            "Bounded timeouts, last-valid policy caching and upstream isolation.",
        ],
        "Testing": [
            "Missing/revoked keys, invalid JWTs, denied scopes, forwarding and header-stripping tests.",
        ],
    },
    "TB4": {
        "Design": [
            "Rate-rule precedence, direct-peer IP policy and RE2 threat-matching rules.",
        ],
        "Coding": [
            "Redis fixed-window, sliding-window and token-bucket enforcement.",
            "Exact IP blocking, threat detection and automatic blocking.",
        ],
        "Non-Functional Requirements": [
            "Atomic shared counters, fail-closed behaviour and race-safe enforcement.",
        ],
        "Testing": [
            "429 limits, consumer isolation, threat, IP-block and Go race tests.",
        ],
    },
    "TB5": {
        "Design": [
            "REST bootstrap, authenticated gRPC updates and sanitised analytics flow.",
        ],
        "Coding": [
            "Policy streaming, reconnect/backoff and bounded asynchronous event reporting.",
            "Analytics APIs and role-aware React dashboard pages.",
        ],
        "Non-Functional Requirements": [
            "Last-valid policy resilience, data minimisation and memory-only browser tokens.",
        ],
        "Testing": [
            "Internal authentication, stream fallback, event queue and dashboard route tests.",
        ],
    },
    "TB6": {
        "Design": [
            "Full-stack Compose topology, release boundary and repeatable demo plan.",
        ],
        "Coding": [
            "Docker Compose wiring, generated secrets and idempotent bootstrap.",
            "Isolated end-to-end script and localhost attack/defence console.",
        ],
        "Non-Functional Requirements": [
            "Reproducible local deployment, safe secret handling and graceful shutdown.",
        ],
        "Testing": [
            "Complete quality gate, allow/deny/scope/rate POC and isolated Docker E2E.",
        ],
    },
}


TESTING_AND_RETRO_DETAIL = {
    "TB1": {
        "testing_detail": "The authentication checks were divided by boundary rather than by endpoint name. Schema cases confirmed that malformed email and missing password data stop before the service. Service cases checked password verification, disabled accounts and the refresh-token lifecycle. Route cases checked that a valid token is still insufficient when the caller lacks the required role. Database-backed cases confirmed that a role change or soft deactivation is reflected on the next protected request. This layered approach matters because a single successful login case cannot show that registration, token renewal, logout and administration fail safely. The focused scenario run also exercised the typed login response and the Admin-only user-management path independently from the broader suite.",
        "test_rows": [
            ("Registration role", "Request Admin through public registration", "Rejected; consumer role retained", "Client JSON cannot escalate privilege"),
            ("Invalid login", "Wrong password for an existing user", "401 structured response", "No session token is issued"),
            ("Refresh lifecycle", "Refresh, revoke, then reuse token", "First succeeds; reused token fails", "Revocation state is enforced"),
            ("Admin boundary", "Consumer calls user-management route", "403 response", "Authentication does not imply administration"),
            ("Valid Admin", "Typed login followed by protected request", "200 response with current role", "Identity and RBAC cooperate"),
        ],
        "retro_detail": "The increment exposed two different meanings of validation. Pydantic can prove that a role string has an allowed shape, but it cannot prove that the current caller may assign that role. Moving privilege decisions into the service and dependency layer removed that ambiguity. The team also learned to test the negative lifecycle, not only token creation: logout is useful only when a previously valid refresh token is demonstrably unusable afterwards. These lessons established the pattern used in later timeboxes: validate shape first, authorise intent second and persist only after both checks succeed.",
        "retro_rows": [
            ("What worked", "Router/service/repository separation made failed branches easy to isolate"),
            ("What changed", "Role assignment became an explicit server-side rule"),
            ("Carry-forward", "Every later administrative resource reused typed input plus RBAC dependencies"),
        ],
    },
    "TB2": {
        "testing_detail": "API-key tests concentrated on disclosure, ownership and lifecycle. Creation had to return a usable raw value once while every later representation remained masked. Repository checks confirmed that only a bcrypt hash and non-sensitive prefix are persisted. Ownership scenarios compared an API Consumer managing its own key with attempts to read or revoke another consumer's key; Admin operations were checked separately. Policy tests exercised field ranges, algorithm enums, referenced key identifiers, duplicate active route permissions and soft-disable behaviour. A focused four-scenario run covered one-time key return and the route-permission validation that prevents two active records from claiming the same method, path and required scope.",
        "test_rows": [
            ("One-time secret", "Create a new API key", "201 with raw key in create response", "Caller receives the secret exactly once"),
            ("Persistence", "Read the stored key record", "Hash and prefix only", "Database disclosure does not reveal the raw key"),
            ("Ownership", "Consumer requests another owner's key", "Controlled denial", "Tenant boundary is applied in the service"),
            ("Duplicate route", "Create the same active method/path/scope", "Validation or conflict response", "Policy selection remains unambiguous"),
            ("Revocation", "Use a key after soft revocation", "403 API_KEY_REVOKED", "History remains, authentication stops"),
        ],
        "retro_detail": "The most important correction was separating permission vocabulary from request routing. Early examples made echo:read look as if the proxy could discover an echo service from the word echo. Testing showed that a label cannot safely define a route, because names can be reused and do not contain an HTTP method or path. RoutePermission therefore became the authoritative mapping, while the key merely carries labels. The timebox also confirmed that list and audit usability should never require redisplaying the raw secret; prefix, name, owner, status, scopes and timestamps are sufficient for administration.",
        "retro_rows": [
            ("What worked", "One-way persistence and masked list responses reduced secret exposure"),
            ("What changed", "RoutePermission became the explicit request-to-scope binding"),
            ("Carry-forward", "Proxy tests used real method/path records rather than interpreting labels"),
        ],
    },
    "TB3": {
        "testing_detail": "The proxy suite used both isolated collaborators and a real upstream integration. Table-driven cases supplied missing, malformed, expired and revoked credentials and asserted that the upstream counter remained unchanged. JWT cases pinned HS256 and rejected an algorithm-substitution attempt rather than trusting the token header. Scope cases compared exact read and write labels, including labels that resemble wildcards; no implicit permission was granted. The forwarding integration then supplied a valid key, active JWT policy and matching RoutePermission, confirmed a 200 upstream response and checked that Aegis credentials were absent from the request received by the backend. Race-enabled tests exercised the same handler code with concurrent requests.",
        "test_rows": [
            ("Missing key", "Request has no X-API-Key", "401 API_KEY_MISSING", "Request stops before upstream"),
            ("Revoked key", "Validated record has revoked state", "403 API_KEY_REVOKED", "Cached metadata cannot bypass lifecycle"),
            ("JWT algorithm", "Token header substitutes another algorithm", "401 JWT error", "Configured algorithm is pinned"),
            ("Exact scope", "GET requires echo:read and key has it", "200 forwarded", "Route and key labels match exactly"),
            ("Denied write", "POST requires echo:write; key has read only", "403 INSUFFICIENT_SCOPE", "Backend is not called"),
            ("Header stripping", "Valid request reaches echo upstream", "Aegis credentials absent", "Secrets do not cross the trust boundary"),
        ],
        "retro_detail": "This increment made ordering visible as a security property. Returning the right status is not enough if a denied request has already reached the backend, so tests observed the upstream call count and received headers. Interfaces around key validation, policy lookup, rate storage and reporting made these assertions possible without a full deployment. The team also learned that cache behaviour belongs in authorisation testing: positive entries need bounded lifetimes, negative entries reduce repeated invalid-key work, and revocation must be reflected quickly enough that convenience never becomes a bypass.",
        "retro_rows": [
            ("What worked", "Interface-driven tests observed both response and upstream side effects"),
            ("What changed", "Exact method/path/scope matching replaced permissive interpretation"),
            ("Carry-forward", "Later controls preserved the same stop-before-upstream invariant"),
        ],
    },
    "TB4": {
        "testing_detail": "Rate-limit verification covered algorithm correctness and distribution semantics. Fixed-window cases checked first use, exhaustion and reset. Sliding-window cases checked that timestamps outside the active interval stop contributing. Token-bucket cases checked burst capacity, refill and rejection when no token remains. Separate tests confirmed that API-key scope isolates consumers while route and global scopes deliberately share counters across proxy instances. Threat tests compiled valid RE2 expressions, rejected unsupported patterns and matched the request target rather than reading a body. IP tests canonicalised the direct peer and applied exact matches. The isolated stack combined these controls and produced the observed 200, 200, 429 sequence plus separate 403 threat and IP outcomes.",
        "test_rows": [
            ("Fixed window", "Three requests under a two-request rule", "200, 200, 429", "Atomic counter enforces the configured limit"),
            ("Key isolation", "Second consumer requests after first is limited", "Second consumer receives 200", "API-key buckets are independent"),
            ("Shared route", "Different keys use one route-scoped rule", "Counter is shared", "Scope semantics match the policy model"),
            ("Threat match", "Request target matches active RE2 rule", "403 THREAT_DETECTED", "Credential work and upstream call are skipped"),
            ("Exact IP", "Direct peer equals an active block", "403 IP_BLOCKED", "Canonical direct-peer policy is enforced"),
            ("Automatic block", "Repeated qualifying violations", "Active block appears in Control Plane", "Protection persists beyond one process"),
        ],
        "retro_detail": "The stale loopback block was a useful failure because it demonstrated that a correct security control can make a test appear broken. The response was not to weaken IP enforcement, but to make environment state explicit: disposable data, unique rule names, pre-run inspection and cleanup. Rate tests also clarified that the word distributed is incomplete without a sharing rule. The test plan now states whether a bucket belongs to one key, one route or the whole proxy, so an observed 429 can be interpreted against the intended scope rather than treated as a generic success.",
        "retro_rows": [
            ("What worked", "Lua scripts made counter updates atomic across instances"),
            ("What changed", "Setup and cleanup became part of security-test repeatability"),
            ("Carry-forward", "Integrated runs record both policy scope and environment state"),
        ],
    },
    "TB5": {
        "testing_detail": "Distribution tests covered the transition from no policy to a usable bootstrap snapshot and then to live gRPC updates. A wrong internal token was rejected. A valid subscriber received a complete snapshot and swapped it without exposing a partially updated policy to concurrent requests. Reconnect cases preserved the last valid snapshot during interruption, while a startup with no valid policy remained closed. Reporter tests filled, flushed and shut down the bounded queue and asserted that delivered events contained decisions but no raw credentials or bodies. Six dashboard tests checked server rendering, role-gated navigation, memory-only authentication state, Admin-only IP and threat pages, and analytics backed by live Control Plane data.",
        "test_rows": [
            ("Bootstrap", "Proxy starts before first stream update", "REST snapshot becomes active", "Initial policy is available deterministically"),
            ("Stream auth", "Subscriber sends wrong internal token", "Connection rejected", "Administrative policy is not publicly distributable"),
            ("Atomic update", "Complete gRPC snapshot arrives", "New policy replaces old snapshot", "Readers never see a partial configuration"),
            ("Interruption", "Control Plane stream disconnects", "Last valid snapshot remains active", "Temporary outage does not remove protection"),
            ("Sanitisation", "Blocked request is reported", "No key, JWT or body in event", "Analytics does not become a secret store"),
            ("Dashboard roles", "Viewer and Consumer navigate protected pages", "Routes follow permission matrix", "Browser controls reflect server roles"),
        ],
        "retro_detail": "The principal lesson was that resilience and freshness pull in opposite directions. Discarding policy on every disconnect would be fresh but unsafe for availability; keeping data forever would be available but potentially stale. The implemented compromise is authenticated complete snapshots, last-valid operation and reconnect with backoff, with operators able to observe the Control Plane state. On the frontend, memory-only tokens avoided persistent browser storage. That choice was retained even though it makes reloads less convenient, because the project prioritises administrative credential protection over uninterrupted browser sessions.",
        "retro_rows": [
            ("What worked", "Whole-snapshot replacement simplified concurrency and rollback reasoning"),
            ("What changed", "Last-valid fallback gained explicit startup and interruption tests"),
            ("Carry-forward", "TB6 exercised policy delivery before every integrated allow/deny case"),
        ],
    },
    "TB6": {
        "testing_detail": "The release check joined component and system-level results without treating one as a substitute for another. make check completed the Python suite with 47 passing tests and 77% application coverage, Go formatting, vet and race checks with 83.4% proxy-core coverage, six dashboard tests and thirteen POC/tooling tests. The disposable Docker scenario then verified live service health, dashboard readiness, Admin login, API-key creation and revocation, JWT and threat policy activation, authenticated gRPC delivery, a successful upstream request, invalid and revoked credentials, manual and automatic exact-IP blocking, the Redis limit sequence and analytics aggregation. Generated credentials were excluded from the retained output and the isolated stack was removed after completion.",
        "test_rows": [
            ("Quality gate", "Run make check", "All component suites and static checks pass", "Repository is internally consistent"),
            ("Service health", "Start disposable Compose stack", "Control Plane, PostgreSQL and Redis healthy", "Dependencies are reachable"),
            ("Policy path", "Activate JWT and subscribe through gRPC", "Valid request reaches upstream with 200", "Management-to-enforcement path works"),
            ("Deny matrix", "Invalid/revoked key, threat and manual block", "Controlled 401 or 403 responses", "Distinct controls stop requests"),
            ("Rate path", "Apply two-request fixed window", "200, 200, 429", "Redis policy works in the deployed stack"),
            ("Analytics", "Query after allow and deny scenarios", "Forwarded and blocked outcomes aggregated", "Reporter and persistence path works"),
            ("Cleanup", "Scenario completes", "Disposable services and data removed", "Run can be repeated safely"),
        ],
        "retro_detail": "The integrated run changed the meaning of done. A component can pass its own tests while deployment wiring, internal authentication, policy propagation or stale state still prevents the product from working. TB6 therefore kept component checks for fast diagnosis and added a bounded end-to-end path for interaction risk. The console remained localhost-only and generated secrets stayed in memory or temporary environment state. Remaining work was recorded honestly: the run demonstrates functional integration, not production throughput, hostile-internet exposure or an independent penetration assessment.",
        "retro_rows": [
            ("What worked", "Layered checks separated component defects from integration defects"),
            ("What changed", "Disposable credentials and cleanup became mandatory release controls"),
            ("Carry-forward", "Future work targets load, accessibility and independent security review"),
        ],
    },
}


FAILURE_AND_RETRO_NOTES = {
    "TB1": {
        "failure_analysis": "The negative cases also distinguish caller mistakes from server faults. Missing fields and invalid formats are client-contract failures; wrong credentials are authentication failures; a valid identity with the wrong role is an authorisation failure. Keeping those outcomes separate gives the dashboard useful feedback without revealing whether a protected account exists. It also makes regression diagnosis faster: a change from 403 to 500 points to exception handling, while an unexpected 200 indicates a security boundary failure. The tests assert both status and response shape so later refactoring cannot silently replace structured errors with framework traces.",
        "next_step": "The carry-forward decision was to require the same three-level check for every new Control Plane resource: typed schema tests, service-level ownership or role tests, and route-level response tests. TB2 adopted this immediately for API-key ownership and policy administration. Session identifiers and user roles were also treated as security data in logs, so debugging output was kept descriptive but did not include token contents.",
    },
    "TB2": {
        "failure_analysis": "A particularly important failure case is a successful database write followed by an unsafe response. The API-key tests therefore inspect output as well as storage: the create response may contain the raw key, but list, read and revoke responses must not. Policy validation follows the inverse concern: a rejected command must not leave a partial active record. Duplicate-route and invalid-reference cases are checked around the service transaction so the system cannot publish ambiguous policy after returning an error to the caller. These assertions protect confidentiality and configuration integrity at the same time.",
        "next_step": "The timebox closed with a clearer documentation rule: examples must always show the scope label beside its RoutePermission, never as an isolated string that appears to discover a backend automatically. TB3 then reused the same example pair in handler tests. Operator pages were designed around masked prefixes and lifecycle state, and any future import or bulk-create feature must preserve one-time disclosure rather than inventing a convenient way to reveal stored secrets.",
    },
    "TB3": {
        "failure_analysis": "Handler tests treat an upstream call during a denied case as a test failure even when the client receives the expected status. This catches implementations that validate too late or write an error after forwarding has begun. The same principle applies to headers: a 200 response is insufficient if the upstream receives X-API-Key or Authorization. By checking side effects, the suite verifies the trust boundary rather than merely the response surface. Timeout and unavailable-Control-Plane cases further confirm that cache and policy behaviour fail in a controlled direction instead of allowing a request because validation could not complete.",
        "next_step": "The retrospective converted the handler order into a review checklist used in TB4 and TB6. Every new control must identify its position, its cost, the status and code it returns, whether it reports an event, and whether the upstream remains untouched. The Go interfaces were kept narrow so new limit, threat and block cases could reuse the same observable call counters rather than require fragile network-only tests for every branch.",
    },
    "TB4": {
        "failure_analysis": "Rate-limit assertions include the sequence and the selected scope because 429 alone can hide an incorrect bucket. If consumer two is limited by consumer one's key-scoped traffic, the algorithm may be atomic but the key design is wrong. Conversely, route and global tests expect deliberate sharing. Threat and IP checks assert their early position by using requests with missing credentials: the expected result remains THREAT_DETECTED or IP_BLOCKED, proving that the request was rejected before key work. Invalid patterns and malformed addresses are administrative validation failures and must not reach the active snapshot.",
        "next_step": "The timebox added an environment preflight to the integrated script: inspect active blocks and rules, create uniquely named demonstration records, use disposable ports and clean up after the run. Test descriptions now include algorithm, scope and window rather than saying only rate limit. A future CIDR feature will require a separate model and precedence design; it must not be simulated by loosening the exact-address matcher that current tests define.",
    },
    "TB5": {
        "failure_analysis": "Policy resilience tests cover both sides of the availability boundary. A subscriber that has already obtained a valid snapshot may keep using it during interruption, but a fresh proxy with no trusted snapshot must not invent empty policy and forward traffic. Event tests similarly distinguish dropping non-critical analytics under queue pressure from dropping the security decision itself: request enforcement completes first, while the bounded reporter handles delivery asynchronously. Dashboard tests remain supporting checks; server-side RBAC is still tested independently because hidden navigation is not an authorisation control.",
        "next_step": "TB5 closed with explicit operational indicators for the final integrated run: service health, snapshot delivery, last update state, reporter completion and analytics visibility. The frontend retained memory-only authentication and documented the reload trade-off. Any future persistent session design will require a fresh threat review covering browser storage, cross-site scripting and logout invalidation rather than being treated as a cosmetic convenience.",
    },
    "TB6": {
        "failure_analysis": "The integrated scenario intentionally repeats controls already covered by unit tests because configuration can invalidate otherwise correct code. It checks that the proxy uses the expected Control Plane address, gRPC token, Redis database, upstream target and generated policy identifiers. Each checkpoint narrows the fault domain: health precedes login, login precedes policy creation, policy delivery precedes the first 200, and denial cases precede analytics inspection. Cleanup is asserted as part of the scenario so a successful run cannot leave active blocks or counters that corrupt the next demonstration.",
        "next_step": "The final retrospective separated release confidence from production assurance. The completed checks support functional correctness in the local architecture, but do not claim capacity under sustained load, resistance to an independent attacker or usability across assistive technologies. Those items were placed in future work with measurable entry criteria: controlled hardware and traffic profiles, an external security review, browser accessibility testing and trusted-load-balancer address handling before any public deployment.",
    },
}


def visible_text(element):
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def find_body_element(doc, exact):
    for element in doc._element.body:
        if visible_text(element) == exact:
            return element
    raise RuntimeError(f"Document boundary not found: {exact}")


def remove_range(doc, first, last, include_first=False):
    body = doc._element.body
    start = body.index(first) + (0 if include_first else 1)
    end = body.index(last)
    for element in list(body)[start:end]:
        body.remove(element)


def move_new_elements_before(doc, start_index, anchor):
    body = doc._element.body
    sect_index = body.index(body.sectPr)
    new_elements = list(body)[start_index:sect_index]
    for element in new_elements:
        body.insert(body.index(anchor), element)


def chapter_heading(doc, title):
    p = helpers.heading(doc, title, 1)
    p.paragraph_format.page_break_before = True
    return p


def compact_table(doc, headers, rows, widths, caption=None):
    return helpers.table(doc, headers, rows, widths, caption)


def add_screenshot_slot(doc, slot, requested_view, capture_instruction):
    compact_table(
        doc,
        [f"SCREENSHOT TO ADD - {slot}: {requested_view}"],
        [],
        [1.0],
    )
    p = helpers.para(doc, f"Capture instruction: {capture_instruction}", italic=True,
                     align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_after = Pt(7)


def add_chapter6(doc):
    chapter_heading(doc, "Chapter 6: System Design")
    helpers.heading(doc, "6.1 Design Basis", 2)
    helpers.para(doc, "The design in this chapter reflects the current repository, database migrations and proxy handler rather than a generic gateway template. Aegis separates configuration from enforcement: operators manage durable identity and policy through the Control Plane and dashboard, while the Go Reverse Proxy makes a bounded decision for each request before any protected upstream is reached. PostgreSQL stores users, keys, policy and events; Redis supplies shared runtime counters and revocation state.")

    helpers.heading(doc, "6.2 Current Architecture", 2)
    helpers.add_figure(doc, DIAGRAMS / "figure_6_1_current_architecture.png", "Figure 6.1: Current Aegis architecture and deployed component connections.", 6.1)
    helpers.para(doc, "The dashboard calls versioned Control Plane REST endpoints. The Control Plane persists administration data in PostgreSQL and uses Redis for runtime state. The proxy first bootstraps policy by REST, then receives authenticated gRPC snapshots. Requests enter only through the proxy; an allowed request is stripped of Aegis credentials and forwarded to the configured upstream. Security outcomes are reported asynchronously so analytics work does not delay the enforcement decision.")
    compact_table(doc, ["Component", "Primary responsibility", "Trust boundary"], [
        ("Dashboard", "Role-aware administration and analytics UI", "Never stores raw issued keys or tokens persistently"),
        ("Control Plane", "Validation, RBAC, policy lifecycle and analytics", "Authoritative administrative API"),
        ("Reverse Proxy", "Per-request authentication, authorisation and traffic controls", "Only allowed path to protected upstreams"),
        ("PostgreSQL", "Durable users, key metadata, policies, blocks and events", "Restricted to the Control Plane"),
        ("Redis", "Revocation state and atomic distributed rate counters", "Internal runtime dependency"),
    ], [1.25, 3.35, 1.9], "Table 6.1: Responsibilities and boundaries in the implemented system.")

    helpers.heading(doc, "6.3 Domain and Persistence Design", 2)
    helpers.add_figure(doc, DIAGRAMS / "figure_6_2_domain_model.png", "Figure 6.2: Persisted domain model derived from current SQLAlchemy models and migrations.", 6.1)
    helpers.para(doc, "User owns APIKey records; the key stores a hash and lookup prefix rather than the raw secret. RateLimitRule may target one key, one route or the whole proxy. JWTConfig stores encrypted signing material and an active state. RoutePermission binds method and path to one exact required scope. ThreatRule and IPBlock represent early rejection policy. SecurityEvent stores sanitised request metadata and the final action for analytics and automatic-block calculations.")
    compact_table(doc, ["Model", "Important fields", "Reason for the structure"], [
        ("APIKey", "owner_id, key_hash, prefix, scopes, revoked_at", "One-time disclosure, fast lookup and traceable revocation"),
        ("RoutePermission", "method, path_pattern, required_scope, is_active", "Makes resource:action labels enforceable against real requests"),
        ("RateLimitRule", "scope_type, target, algorithm, limit, window", "Supports specific-to-general rule precedence"),
        ("SecurityEvent", "request_id, action, reason, status, route, client_ip", "Supports audit and analytics without storing credentials or bodies"),
    ], [1.3, 2.4, 2.8], "Table 6.2: Selected design choices in the persisted model.")

    helpers.heading(doc, "6.4 Administrative Request Validation", 2)
    helpers.add_figure(doc, DIAGRAMS / "figure_6_4_control_plane_validation.png", "Figure 6.3: Control Plane validation from JSON decoding to persistence.", 6.1)
    helpers.para(doc, "FastAPI decodes JSON and asks a typed Pydantic schema to validate required fields, enum values, address formats and numeric ranges. Authentication and RBAC dependencies run before the service. The service then checks rules that depend on existing data, such as ownership, duplicate active policy and referenced API keys. Only a valid command reaches the repository transaction. This division matters: field validation returns a controlled 422 response, authorisation failures return 401 or 403, and domain conflicts are handled without partial writes.")
    helpers.code(doc, '''{
  "method": "GET",
  "path_pattern": "/api/echo",
  "required_scope": "echo:read",
  "is_active": true
}''', "Example: a route-permission object binds the concrete GET request to echo:read.")

    helpers.heading(doc, "6.5 Proxy Enforcement Sequence", 2)
    helpers.add_figure(doc, DIAGRAMS / "figure_6_3_enforcement_sequence.png", "Figure 6.4: Core request-validation and enforcement sequence in the Go proxy.", 6.1)
    helpers.para(doc, "The handler creates a request identifier and canonical client address, rejects an exact IP block or matching threat, validates the API key, validates JWT policy when active, resolves the matching route permission, compares the exact required scope and consumes the selected Redis limit. A rejection returns immediately; it does not call the upstream. An allowed request has Aegis credentials removed before ReverseProxy.ServeHTTP forwards it. The final action is sent to the asynchronous reporter.")

    helpers.heading(doc, "6.6 Scope and Resource Mapping", 2)
    helpers.para(doc, "A label such as echo:read consists of a resource name and action separated by a colon. It remains a string because keys may carry several independent permissions. The proxy does not infer what echo means. It finds the active RoutePermission whose method and path pattern match the incoming request, reads required_scope, and performs an exact membership check against the validated key scopes. Therefore [\"echo:read\"] is useful only when a route record explicitly maps a real request to echo:read.")
    compact_table(doc, ["Incoming request", "Route permission", "Key scopes", "Decision"], [
        ("GET /api/echo", "echo:read", "[echo:read]", "Allow"),
        ("POST /api/echo", "echo:write", "[echo:read]", "Deny: insufficient scope"),
        ("POST /api/echo", "echo:write", "[echo:read, echo:write]", "Allow"),
        ("GET /api/orders", "orders:read", "[echo:read]", "Deny: insufficient scope"),
    ], [1.55, 1.35, 1.9, 1.7], "Table 6.3: Exact route-to-scope examples.")

    helpers.heading(doc, "6.7 Policy Distribution and Resilience", 2)
    helpers.add_figure(doc, DIAGRAMS / "figure_6_5_policy_distribution.png", "Figure 6.5: REST bootstrap, authenticated gRPC updates and last-valid snapshot behaviour.", 6.1)
    helpers.para(doc, "At startup the proxy retrieves a complete policy snapshot through the internal REST contract. It then subscribes to gRPC updates authenticated with a separate internal token. A complete new snapshot replaces the previous in-memory value atomically. If the stream disconnects, reconnect logic runs with backoff while the last valid policy continues to protect traffic. The proxy fails closed when no initial policy is available.")

    helpers.heading(doc, "6.8 Rate-Limit Decision", 2)
    helpers.add_figure(doc, DIAGRAMS / "figure_6_6_rate_limit_decision.png", "Figure 6.6: Rate-limit rule selection and Redis-backed decision flow.", 6.1)
    helpers.para(doc, "The selector evaluates active rules from most specific to least specific. API-key rules isolate each consumer; route rules share capacity for the matching route; global rules share one bucket. Redis Lua scripts make reads and writes atomic across proxy instances. Fixed window is simple, sliding window smooths boundary bursts, and token bucket supports controlled bursts with refill. A denied request returns 429 and reset metadata.")

    helpers.heading(doc, "6.9 Event and Analytics Flow", 2)
    helpers.add_figure(doc, DIAGRAMS / "figure_6_7_event_analytics.png", "Figure 6.7: Sanitised event reporting, analytics aggregation and automatic IP blocking.", 6.1)
    helpers.para(doc, "The proxy reports request identifier, route, method, client address, status, action, reason and timing, but excludes raw keys, JWTs and bodies. A bounded queue prevents reporting from blocking requests. The Control Plane persists the event, aggregates dashboard metrics and may create an automatic exact-IP block after repeated qualifying violations. This feedback path is deliberately separate from the immediate decision path.")

    helpers.heading(doc, "6.10 Security Trade-offs and Summary", 2)
    helpers.para(doc, "Aegis chooses exact routes, exact scopes and direct-peer IP values to avoid silently trusting ambiguous wildcard or forwarded-header input. These choices are safer for the current deployment but limit CIDR policy and operation behind trusted load balancers. Memory-only dashboard tokens reduce browser persistence but require login after reload. Last-valid policy improves availability, while authenticated distribution and fail-closed startup prevent an untrusted or empty snapshot from weakening enforcement. The result is a clear boundary: the Control Plane defines valid policy, and the proxy applies it consistently before upstream access.")


def add_chapter7(doc):
    chapter_heading(doc, "Chapter 7: Project Timeline")
    helpers.heading(doc, "7.1 Planning Basis", 2)
    helpers.para(doc, "The Project Proposal fixes Design between 27 April and 29 May 2026, Implementation between 1 June and 11 September 2026, and the separate project Testing activity between 14 September and 9 October 2026 (Htun Khaing Lynn, 2026). Only Implementation is divided into timeboxes. Design remains an enabling activity before timeboxed delivery, while the formal Testing and Evaluation periods remain ordinary project activities after implementation. This avoids incorrectly calling every project phase a timebox.")
    helpers.para(doc, "The proposal specifies six implementation timeboxes and allocates seventy-five working days to Implementation. The schedule therefore retains six fixed timeboxes and distributes the complete implementation window across them. Each timebox is approximately two and a half weeks, within the usual DSDM range of two to four weeks. Weekends are excluded from the effort calculation, while the calendar start and finish dates remain fixed.")

    helpers.heading(doc, "7.2 DSDM Structured Timebox Rules", 2)
    helpers.para(doc, "Aegis uses the DSDM structured timebox rather than a Scrum phase breakdown. For readability, the report uses plain implementation labels for the five official DSDM control points: Implementation Setup corresponds to Kick-off, Feature Build to Investigation, Testing and Fixing to Refinement, Integration and Documentation to Consolidation, and Review and Sign-off to Close-out. The labels are simplified, but their sequence, review purpose and fixed-time behaviour remain unchanged (Agile Business Consortium, 2014).")
    compact_table(doc, ["Timebox step", "12-day TB", "13-day TB", "Aegis application"], [
        ("Implementation Setup", "0.5 day", "0.5 day", "Confirm objective, availability, dependencies, MoSCoW scope and acceptance criteria."),
        ("Feature Build", "2 days", "2 days", "Clarify detailed requirements and technical approach, then build the selected features."),
        ("Testing and Fixing", "7 days", "8 days", "Test iteratively, correct defects and demonstrate the near-complete increment."),
        ("Integration and Documentation", "2 days", "2 days", "Integrate the increment, complete regression and security checks, and update documentation."),
        ("Review and Sign-off", "0.5 day", "0.5 day", "Accept the increment, record the outcome and agree actions for the next timebox."),
    ], [1.55, 0.85, 0.85, 3.25], "Table 7.1: Day allocation for the 12-day and 13-day Aegis implementation timeboxes.")
    helpers.para(doc, "A 12-day timebox totals 0.5 + 2 + 7 + 2 + 0.5 days, while a 13-day timebox totals 0.5 + 2 + 8 + 2 + 0.5 days. Testing remains integrated throughout the main delivery work and is followed by explicit integration, documentation and acceptance activity; it is not postponed until the separate project Testing period.")

    helpers.heading(doc, "7.3 MoSCoW Capacity Rule", 2)
    helpers.para(doc, "Time and quality are fixed; detailed scope is the variable. Each timebox reserves approximately 55% of effort for Must Haves, 25% for Should Haves and 20% for Could Haves. Won't Haves are excluded from capacity. This keeps Must Have effort below the DSDM recommendation of 60% and maintains a real 20% contingency. A Must Have is used only where the objective would fail without it; a Should Have remains important but has a workaround; a Could Have is the first scope removed when risk materialises.")
    compact_table(doc, ["Priority", "Effort", "Meaning in Aegis"], [
        ("Must Have", "55%", "Essential to the usable security increment; without it the timebox objective is not met."),
        ("Should Have", "25%", "Important and expected, but the increment remains usable with a documented workaround."),
        ("Could Have", "20%", "Contingency pool delivered only after Must/Should quality remains protected."),
        ("Won't Have this time", "Excluded", "Explicitly deferred so it cannot silently consume the fixed timebox."),
    ], [1.5, 1.0, 4.0], "Table 7.2: MoSCoW effort balance used for every implementation timebox.")

    helpers.heading(doc, "7.4 Implementation-Only Delivery Plan", 2)
    compact_table(doc, ["ID", "Implementation increment", "Fixed dates", "Working days"], [
        ("TB1", "Control Plane Foundation and Identity", "1-17 Jun", "13"),
        ("TB2", "API Keys and Administrative Policy", "18 Jun-3 Jul", "12"),
        ("TB3", "Proxy Authentication and Authorisation", "6-22 Jul", "13"),
        ("TB4", "Distributed Limits and Threat Enforcement", "23 Jul-7 Aug", "12"),
        ("TB5", "Policy Distribution, Analytics and Dashboard", "10-26 Aug", "13"),
        ("TB6", "Integration, POC and Hardening", "27 Aug-11 Sep", "12"),
    ], [0.6, 2.9, 1.8, 1.0], "Table 7.3: Six fixed timeboxes across the proposal's 75-working-day implementation window.")
    helpers.add_figure(doc, DIAGRAMS / "aegis_dsdm_implementation_timeline.png", "Figure 7.1: Design, six implementation timeboxes and the separate project-testing period.", 6.1)

    for idx, tb in enumerate(helpers.TIMEBOXES, 1):
        helpers.page_break(doc)
        helpers.heading(doc, f"7.{4 + idx} {tb['id']} - {tb['name']}", 2)
        start_date, end_date = [part.strip() for part in tb["dates"].split(" - ", 1)]
        compact_table(doc, ["Timebox Detail", "Value"], [
            ("Timebox Name", f"{tb['name']} Timebox"),
            ("Start Date", start_date),
            ("End Date", end_date),
        ], [1.55, 4.95])
        compact_table(
            doc,
            ["Task", "Duration", "Start Date", "End Date"],
            TIMEBOX_SIMPLE_SCHEDULES[tb["id"]],
            [2.65, 0.85, 1.5, 1.5],
            f"Table 7.{idx + 3}: {tb['id']} task and date schedule.",
        )

        helpers.heading(doc, "Key Deliverables (Output)", 3)
        for group, items in TIMEBOX_SIMPLE_DELIVERABLES[tb["id"]].items():
            p = helpers.para(doc, group, bold_lead=group, align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.keep_with_next = True
            helpers.bullets(doc, items)

    outside_heading = helpers.heading(doc, "7.11 Activities Outside the Timeboxes", 2)
    outside_heading.paragraph_format.page_break_before = True
    compact_table(doc, ["Activity", "Dates", "Relationship to timeboxing"], [
        ("Design", "27 Apr-29 May 2026", "Enough Design Up Front: architecture, data, API and security foundations; not a timebox."),
        ("Implementation", "1 Jun-11 Sep 2026", "The only activity divided into six DSDM structured timeboxes."),
        ("Project Testing", "14 Sep-9 Oct 2026", "Independent consolidation of unit, integration, security, performance and UAT results; not a timebox."),
        ("Evaluation and conclusion", "12-30 Oct 2026", "Evaluation against objectives and reflection; not a timebox."),
    ], [1.6, 1.6, 3.3], "Table 7.10: Separation between implementation timeboxes and non-timeboxed project activities.")

    helpers.heading(doc, "7.12 Chapter Summary", 2)
    helpers.para(doc, "The plan keeps six implementation-only timeboxes. Each is presented with a simple name/date block, a task schedule with duration and a concise list of outputs. The five DSDM control points, MoSCoW contingency, integrated testing and fixed-date sign-off are retained without repeating detailed definition or acceptance-boundary text. This protects the 11 September implementation finish. Design remains an enabling activity, while separate project Testing continues to 9 October.")


def add_timebox_examples(doc, tb_id):
    if tb_id == "TB1":
        helpers.code(doc, '''POST /api/v1/auth/login
{
  "email": "admin@example.test",
  "password": "<development-password>"
}''', "Example request: login data is typed and validated before authentication.")
        helpers.add_figure(doc, IMAGES / "figure_8_1_dashboard_login.png", "Figure 8.1: Aegis dashboard login page with memory-only session guidance.", 6.1)
        add_screenshot_slot(doc, "8-A", "Admin user-management page", "Log in as Admin, open Users, show the role and active-state controls, and hide all real email addresses.")
    elif tb_id == "TB2":
        helpers.code(doc, '''POST /api/v1/api-keys
{
  "name": "Echo reader",
  "scopes": ["echo:read"],
  "expires_at": "2026-12-31T23:59:59Z"
}

POST /api/v1/route-permissions
{
  "method": "GET",
  "path_pattern": "/api/echo",
  "required_scope": "echo:read",
  "is_active": true
}''', "Example pair: create a least-privilege key, then bind a real route to its scope.")
        helpers.add_figure(doc, IMAGES / "figure_8_2_openapi_control_plane.png", "Figure 8.2: Generated Control Plane OpenAPI page showing protected and public operations.", 6.1)
        add_screenshot_slot(doc, "8-B", "Dashboard API Keys page", "Create a demonstration key, copy the raw value off-screen, then capture the list showing only name, prefix, scopes, status and expiry.")
    elif tb_id == "TB3":
        helpers.code(doc, '''# Key contains ["echo:read"]
GET  /api/echo  -> required echo:read  -> 200
POST /api/echo  -> required echo:write -> 403 INSUFFICIENT_SCOPE''', "Example decision: the same consumer is allowed to read but not write.")
    elif tb_id == "TB4":
        helpers.code(doc, '''POST /api/v1/rate-limit-rules
{
  "name": "Echo demo limit",
  "scope_type": "global",
  "algorithm": "fixed_window",
  "limit": 2,
  "window_seconds": 60,
  "is_active": true
}

# Observed proxy results
200, 200, 429 RATE_LIMIT_EXCEEDED''', "Example rule and resulting request sequence.")
        add_screenshot_slot(doc, "8-C", "Dashboard Rate Limits and Threat Rules pages", "Use development data only. Capture the active limit and one threat rule without exposing secrets or internal tokens.")
    elif tb_id == "TB5":
        helpers.code(doc, '''{
  "request_id": "req-demo-001",
  "method": "POST",
  "route": "/api/echo",
  "status_code": 403,
  "action": "blocked",
  "reason": "INSUFFICIENT_SCOPE"
}''', "Example sanitised event: it contains a decision but no API key, JWT or body.")
        add_screenshot_slot(doc, "8-D", "Dashboard Analytics page", "After running allow, deny and rate-limit cases, capture the summary cards and recent-event table with test addresses only.")
    elif tb_id == "TB6":
        helpers.code(doc, '''make check
./scripts/e2e.sh

# Integrated outcomes
health -> policy sync -> upstream 200
threat -> 403
rate sequence -> 200, 200, 429''', "Example release commands and the main integrated outcomes.")
        helpers.add_figure(doc, IMAGES / "figure_8_3_attack_console_ui.png", "Figure 8.3: Localhost-only attack/defence console used to run controlled proxy scenarios.", 6.1)


def add_chapter8(doc):
    chapter_heading(doc, "Chapter 8: System Implementation")
    helpers.heading(doc, "8.1 Implementation Approach", 2)
    helpers.para(doc, "Implementation followed six fixed-date increments. To make the chapter easier to review, every timebox uses the same four-part structure: Design states the decision made for the increment; Coding identifies how the repository realises it; Testing records the observed checks and outcomes; Retro states what changed in understanding or process. Examples and interface captures are placed beside the feature that they explain.")

    for tb_index, tb in enumerate(TIMEBOXES, 1):
        sec = tb["section"]
        detail = TESTING_AND_RETRO_DETAIL[tb["id"]]
        notes = FAILURE_AND_RETRO_NOTES[tb["id"]]
        helpers.heading(doc, f"{sec} {tb['id']} - {tb['name']} ({tb['dates']})", 2)
        helpers.heading(doc, f"{sec}.1 Design", 3)
        helpers.para(doc, tb["design"])
        helpers.heading(doc, f"{sec}.2 Coding", 3)
        helpers.para(doc, tb["coding"])
        add_timebox_examples(doc, tb["id"])
        helpers.heading(doc, f"{sec}.3 Testing", 3)
        helpers.para(doc, tb["testing"])
        helpers.para(doc, detail["testing_detail"])
        testing_table_number = 2 * tb_index - 1
        compact_table(
            doc,
            ["Check", "Trigger", "Observed result", "Interpretation"],
            detail["test_rows"],
            [1.15, 1.75, 1.75, 1.85],
            f"Table 8.{testing_table_number}: {tb['id']} testing record.",
        )
        helpers.para(doc, notes["failure_analysis"])
        helpers.heading(doc, f"{sec}.4 Retro", 3)
        helpers.para(doc, tb["retro"])
        helpers.para(doc, detail["retro_detail"])
        compact_table(
            doc,
            ["Retrospective point", "Recorded outcome"],
            detail["retro_rows"],
            [1.6, 4.9],
            f"Table 8.{testing_table_number + 1}: {tb['id']} retrospective actions.",
        )
        helpers.para(doc, notes["next_step"])

    helpers.heading(doc, "8.8 Cross-Component Request Example", 2)
    helpers.para(doc, "Consider GET /api/echo with an active key carrying echo:read. The dashboard or API client first creates a RoutePermission that requires echo:read for that method and path. The Control Plane validates and persists it, the gRPC server publishes a complete snapshot, and the Go subscriber swaps the snapshot into memory. At request time the proxy validates the key and JWT policy, resolves the route, finds the exact scope, consumes the chosen limit, strips credentials and forwards. POST /api/echo follows the same path but requires echo:write and is denied when that label is absent.")
    compact_table(doc, ["Layer", "Input", "Successful output", "Common failure"], [
        ("Pydantic", "Administrative JSON", "Typed command", "422 invalid field or enum"),
        ("Service", "Typed command + current user", "Authorised domain change", "403 role or ownership denial"),
        ("Repository", "Validated domain change", "Committed row", "Controlled duplicate conflict"),
        ("Policy stream", "Complete database snapshot", "Authenticated proxy update", "Reconnect with last-valid snapshot"),
        ("Proxy", "HTTP request + in-memory policy", "Sanitised forwarded request", "401, 403 or 429 without upstream call"),
    ], [1.1, 1.8, 2.0, 2.0], "Table 8.13: One request contract across the Aegis components.")

    helpers.heading(doc, "8.9 Screenshot Completion Guide", 2)
    helpers.para(doc, "The report already includes the login page, generated OpenAPI page and attack/defence console. Four marked slots remain for authenticated pages that require a live signed-in dashboard session. When adding them, use development-only data, crop browser chrome where appropriate, hide emails and addresses that identify real people, and never display a raw API key, bearer token, password or internal service token.")
    compact_table(doc, ["Slot", "Report page", "Page to capture", "What must be visible"], [
        ("8-A", "46", "Users", "Role and active-state controls with anonymised users"),
        ("8-B", "49", "API Keys", "Name, prefix, scopes, status and expiry; no raw key"),
        ("8-C", "53", "Rate Limits / Threat Rules", "One active demonstration rule in each view"),
        ("8-D", "55", "Analytics", "Summary cards and sanitised recent outcomes"),
    ], [0.55, 0.85, 1.65, 3.45], "Table 8.14: Webpage captures still required from an authenticated session.")

    helpers.heading(doc, "8.10 Implementation Summary", 2)
    helpers.para(doc, "The implementation now connects typed administrative JSON, durable policy, authenticated distribution and ordered Go enforcement. The repeated Design-Coding-Testing-Retro format shows how each increment progressed and what was learned. The practical examples explain why a key contains scope strings, how a route gives those strings meaning, how invalid input is rejected, and how distributed policy produces observable allow or deny outcomes without exposing secrets.")


def replace_front_matter(doc):
    toc = find_body_element(doc, "Table of Contents")
    list_figures = find_body_element(doc, "List of Figures")
    remove_range(doc, toc, list_figures)
    body = doc._element.body
    start = body.index(body.sectPr)
    for title, page in TOC_PAGES.items():
        helpers.toc_entry(doc, title, page, 0)
    helpers.page_break(doc)
    move_new_elements_before(doc, start, list_figures)

    list_tables = find_body_element(doc, "List of Tables")
    remove_range(doc, list_figures, list_tables)
    start = body.index(body.sectPr)
    compact_table(doc, ["Figure", "Description"], [
        ("Figure 6.1", "Current architecture"),
        ("Figure 6.2", "Persisted domain model"),
        ("Figure 6.3", "Control Plane request validation"),
        ("Figure 6.4", "Proxy enforcement sequence"),
        ("Figure 6.5", "Policy distribution and fallback"),
        ("Figure 6.6", "Rate-limit decision"),
        ("Figure 6.7", "Event and analytics flow"),
        ("Figure 7.1", "Date-based delivery schedule"),
        ("Figure 8.1", "Dashboard login page"),
        ("Figure 8.2", "Control Plane OpenAPI page"),
        ("Figure 8.3", "Attack/defence console"),
    ], [1.2, 5.3])
    helpers.page_break(doc)
    move_new_elements_before(doc, start, list_tables)

    ch1 = find_body_element(doc, "Chapter 1: Introduction")
    remove_range(doc, list_tables, ch1)
    start = body.index(body.sectPr)
    compact_table(doc, ["Chapter group", "Principal tables"], [
        ("1-5", "SWOT, product comparison, requirements, roles and MoSCoW priorities"),
        ("6", "Component boundaries, data choices and exact route-to-scope examples"),
        ("7", "DSDM allocation, MoSCoW plan and six simple task/date/deliverable timebox records"),
        ("8", "Cross-component request contract and screenshot completion guide"),
        ("9", "Objectives, strengths, limitations and future improvements"),
    ], [1.35, 5.15])
    helpers.page_break(doc)
    move_new_elements_before(doc, start, ch1)


def remove_forbidden_word(doc):
    for node in doc._element.iter(qn("w:t")):
        if node.text:
            node.text = re.sub(r"\bEvidence\b", "Verification", node.text)
            node.text = re.sub(r"\bevidence\b", "verification", node.text)
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            for paragraph in part.paragraphs:
                for run in paragraph.runs:
                    run.text = re.sub(r"\bEvidence\b", "Verification", run.text)
                    run.text = re.sub(r"\bevidence\b", "verification", run.text)


def enforce_arial_black(doc):
    """Keep the requested all-black Arial system, including code examples."""
    for run in doc._element.iter(qn("w:r")):
        r_pr = run.find(qn("w:rPr"))
        if r_pr is None:
            r_pr = OxmlElement("w:rPr")
            run.insert(0, r_pr)
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, fonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), "Arial")
        color = r_pr.find(qn("w:color"))
        if color is None:
            color = OxmlElement("w:color")
            r_pr.append(color)
        color.set(qn("w:val"), "000000")
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            for paragraph in part.paragraphs:
                for run in paragraph.runs:
                    helpers.set_run(run, size=run.font.size.pt if run.font.size else 9, name="Arial")


def build():
    doc = Document(SOURCE)
    body = doc._element.body
    ch6 = find_body_element(doc, "Chapter 6: System Design")
    ch9 = find_body_element(doc, "Chapter 9: Evaluation and Reflection")
    remove_range(doc, ch6, ch9, include_first=True)

    insert_at = body.index(body.sectPr)
    add_chapter6(doc)
    add_chapter7(doc)
    add_chapter8(doc)
    move_new_elements_before(doc, insert_at, ch9)

    replace_front_matter(doc)
    remove_forbidden_word(doc)
    enforce_arial_black(doc)
    doc.core_properties.title = "Aegis API Security and Management System - Compact Project Documentation"
    doc.core_properties.subject = "COMP 1682 project report, compact submission edition"
    doc.core_properties.keywords = "Aegis, API security, reverse proxy, FastAPI, Go, DSDM"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
