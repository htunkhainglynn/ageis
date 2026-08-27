from docx import Document
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "/Users/htunkhainglynn/Projects/aegis/aegis-backend/Aegis_Reverse_Proxy_Reverse_Learning_Handbook.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.18)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True)
        shade(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def p(doc, text):
    para = doc.add_paragraph(text)
    para.paragraph_format.space_after = Pt(6)
    return para


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc, items):
    for index, item in enumerate(items, 1):
        para = doc.add_paragraph(f"{index}. {item}")
        para.paragraph_format.left_indent = Inches(0.2)
        para.paragraph_format.first_line_indent = Inches(-0.2)
        para.paragraph_format.space_after = Pt(3)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)
for name, size, color in [
    ("Title", 24, "17365D"),
    ("Heading 1", 16, "17365D"),
    ("Heading 2", 13, "1F4E79"),
    ("Heading 3", 11, "3B3B3B"),
]:
    st = styles[name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)

doc.add_heading("Aegis Reverse Proxy Reverse-Learning Handbook", 0)
p(doc, "A project-specific beginner guide for learning Go, reverse proxy behavior, JSON payload design, request validation, policy enforcement, and how the pieces in this repository work together.")
p(doc, "Source reviewed: /Users/htunkhainglynn/Projects/aegis/reverse-proxy")
p(doc, "Generated into the backend workspace so Finder can locate it beside your backend project.")
doc.add_page_break()

h(doc, "1. What This Project Is")
p(doc, "This service is the Aegis data plane. It receives public HTTP requests, decides whether each request is allowed, strips credentials, and then forwards the request to the configured backend service. The backend sees a normal request, while the proxy handles security checks at the edge.")
p(doc, "A reverse proxy is a server that sits in front of another server. Clients call the proxy. The proxy calls the backend. In this project the proxy is not only forwarding traffic; it also acts like a security gate.")
bullets(doc, [
    "It validates an API key by calling the Control Plane.",
    "It validates a JWT using policy received from the Control Plane.",
    "It checks whether the client IP is blocked.",
    "It checks method plus URL against threat regex rules.",
    "It enforces rate limits in Redis.",
    "It reports sanitized security events back to the Control Plane.",
])

h(doc, "2. Big Picture Architecture")
add_table(doc, ["Component", "Where", "Purpose"], [
    ("cmd/proxy/main.go", "Application startup", "Loads config, constructs dependencies, starts HTTP server, handles shutdown."),
    ("internal/config", "Config package", "Reads environment variables and validates required settings."),
    ("internal/proxy/handler.go", "Request handler", "Runs the security pipeline and forwards allowed requests."),
    ("internal/controlplane", "HTTP/gRPC clients", "Validates API keys, fetches policy, subscribes to streaming policy updates, reports events."),
    ("internal/proxy/policy.go", "Policy structs", "Defines the JSON-shaped policy snapshot used by validators and guards."),
    ("internal/proxy/rate_limiter.go", "Redis rate limiter", "Enforces fixed-window, sliding-window, and token-bucket limits."),
], [1.7, 1.7, 3.6])

h(doc, "3. End-to-End Request Flow")
numbered(doc, [
    "Client sends a request to the reverse proxy.",
    "The proxy creates a validation context with a timeout so security checks cannot hang forever.",
    "The IP blocker parses r.RemoteAddr and compares it against the policy's blocked_ip_addresses list.",
    "The threat detector builds a target string from request method plus URL and tests it against policy regex patterns.",
    "The handler reads the configured API key header, default X-API-Key.",
    "The Control Plane validator posts {\"api_key\":\"...\"} to the private validation endpoint.",
    "The JWT validator reads Authorization: Bearer <jwt> and validates algorithm, signature, expiration, issuer, and audience.",
    "The Redis rate limiter selects the best matching rule and runs the corresponding Redis script.",
    "If everything passes, the proxy clones the request, removes the API key and Authorization headers, and forwards it.",
    "The event reporter reports the outcome without credentials, bodies, query strings, or arbitrary headers.",
])

h(doc, "4. Why JSON Payloads Have This Shape")
p(doc, "The JSON structures are shaped around ownership boundaries. The proxy does not own API key storage or policy editing. The Control Plane owns those. So JSON becomes the contract between the Control Plane and proxy.")
p(doc, "In Go, JSON tags connect exported struct fields to wire names. For example APIKey becomes api_key because the external API uses snake_case while Go code conventionally uses CamelCase.")

h(doc, "5. API Key Validation Request")
add_code(doc, 'type validationRequest struct {\n    APIKey string `json:"api_key"`\n}')
add_code(doc, '{\n  "api_key": "ak_live_example"\n}')
p(doc, "This request is intentionally tiny because the proxy only needs to prove possession of one secret: the API key presented by the client. It does not send headers, body, route metadata, or JWT data to validate the key. That keeps the Control Plane validation route focused and reduces accidental data exposure.")

h(doc, "6. API Key Validation Response")
add_code(doc, 'type apiResponse struct {\n    Status  string             `json:"status"`\n    Message string             `json:"message"`\n    Data    *proxycore.KeyInfo `json:"data"`\n}\n\ntype KeyInfo struct {\n    ID        int64      `json:"id"`\n    KeyPrefix string     `json:"key_prefix"`\n    OwnerID   int64      `json:"owner_id"`\n    Scopes    []string   `json:"scopes"`\n    Status    string     `json:"status"`\n    ExpiresAt *time.Time `json:"expires_at"`\n}')
p(doc, "The envelope has status, message, and data. The status/message fields are useful for API consistency. The data field contains the actual key metadata the proxy needs. Data is a pointer so the code can detect whether the Control Plane omitted it.")
doc.add_page_break()
add_table(doc, ["Field", "Why It Exists", "How It Is Validated/Used"], [
    ("id", "Stable key identity.", "Stored on security events and used for api_key rate-limit buckets."),
    ("key_prefix", "Human-safe key reference.", "Can identify a key without exposing the full secret."),
    ("owner_id", "Links key to account/project owner.", "Currently carried as metadata for future policy decisions."),
    ("scopes", "Permissions attached to the key.", "Decoded and retained; route-scope enforcement is not implemented in this handler yet."),
    ("status", "Lifecycle state.", "active is allowed, revoked becomes API_KEY_REVOKED, unsupported status is treated as validator unavailable."),
    ("expires_at", "Optional expiration timestamp.", "nil means no expiry; non-nil timestamp must be after time.Now()."),
], [1.15, 2.45, 3.7])

h(doc, "7. Policy Snapshot JSON")
add_code(doc, 'type PolicySnapshot struct {\n    GeneratedAt        time.Time         `json:"generated_at"`\n    JWT                *JWTPolicy        `json:"jwt"`\n    RateLimitRules     []RateLimitPolicy `json:"rate_limit_rules"`\n    BlockedIPAddresses []string          `json:"blocked_ip_addresses"`\n    ThreatRules        []ThreatPolicy    `json:"threat_rules"`\n}')
add_code(doc, '{\n  "generated_at": "2026-07-31T15:30:00Z",\n  "jwt": {"id": 1, "algorithm": "HS256", "verification_key": "secret", "issuer": "aegis", "audience": "api"},\n  "rate_limit_rules": [{"id": 10, "scope_type": "global", "scope_value": "", "algorithm": "fixed_window", "limit_count": 100, "window_seconds": 60, "burst_allowance": null}],\n  "blocked_ip_addresses": ["203.0.113.10"],\n  "threat_rules": [{"id": 7, "name": "SQL injection probe", "pattern": "(?i)(union select|sleep\\\\()", "severity": "high"}]\n}')
p(doc, "This snapshot shape is practical because the proxy needs one coherent read model: JWT policy, rate-limit policy, IP blocks, and threat rules together. The streaming provider can swap in the latest valid snapshot without every validator making its own network call.")

h(doc, "8. JWT Policy")
add_table(doc, ["Field", "Purpose", "Validation Behavior"], [
    ("algorithm", "Allowed signing algorithm: HS256, RS256, or ES256.", "Parser is configured with WithValidMethods and also checks token.Method.Alg()."),
    ("verification_key", "Secret or public key used to verify signature.", "HS256 requires non-empty string; RS256 and ES256 parse PEM public keys."),
    ("issuer", "Optional expected issuer.", "When non-empty, jwt.WithIssuer enforces it."),
    ("audience", "Optional expected audience.", "When non-empty, jwt.WithAudience enforces it."),
], [1.3, 2.5, 3.2])
p(doc, "The JWT validator also requires exp. That matters because a valid signature alone is not enough; without expiration, a leaked token could stay useful forever.")

h(doc, "9. Rate Limit Policy")
add_code(doc, 'type RateLimitPolicy struct {\n    ID             int64  `json:"id"`\n    ScopeType      string `json:"scope_type"`\n    ScopeValue     string `json:"scope_value"`\n    Algorithm      string `json:"algorithm"`\n    LimitCount     int64  `json:"limit_count"`\n    WindowSeconds  int64  `json:"window_seconds"`\n    BurstAllowance *int64 `json:"burst_allowance"`\n}')
add_table(doc, ["Concept", "Accepted Values", "Meaning"], [
    ("scope_type", "api_key, route, global", "Rule matching priority is api_key first, route second, global last."),
    ("scope_value", "key id, path, or empty", "Compared against keyInfo.ID for api_key or exact request path for route."),
    ("algorithm", "fixed_window, sliding_window, token_bucket", "Chooses the Redis Lua script."),
    ("limit_count", "positive integer", "Maximum allowed count in the window."),
    ("window_seconds", "positive integer", "Duration of the enforcement window."),
    ("burst_allowance", "integer or null", "Extra capacity only used by token_bucket."),
], [1.3, 1.9, 3.8])
p(doc, "The code validates semantic rules at enforcement time: limit_count and window_seconds must be positive, algorithm must be supported, Redis must return three result values, and Redis values must be integers.")

h(doc, "10. Threat Rules and IP Blocks")
p(doc, "Threat rules are regex policies. The detector builds target = METHOD + space + request.URL.RequestURI(). That means a rule can match both the HTTP method and path/query string. Invalid regex patterns make policy unavailable because a bad security policy is not safe to silently ignore.")
add_code(doc, 'type ThreatPolicy struct {\n    ID       int64  `json:"id"`\n    Name     string `json:"name"`\n    Pattern  string `json:"pattern"`\n    Severity string `json:"severity"`\n}')
p(doc, "IP blocks are exact IP addresses, not CIDR ranges. The proxy parses RemoteAddr, strips port and brackets, normalizes IPv4-mapped addresses, and compares them with blocked_ip_addresses.")

h(doc, "11. Error Response JSON")
add_code(doc, 'type errorBody struct {\n    Status    string `json:"status"`\n    ErrorCode string `json:"errorCode"`\n    Message   string `json:"message"`\n}')
add_table(doc, ["HTTP", "errorCode", "When It Happens"], [
    ("400", "CLIENT_IP_INVALID", "RemoteAddr cannot be parsed as an IP address."),
    ("401", "API_KEY_MISSING", "Configured API key header is empty."),
    ("401", "API_KEY_INVALID", "Control Plane reports unknown key."),
    ("401", "JWT_MISSING / JWT_INVALID / JWT_EXPIRED", "Authorization bearer token is missing, malformed, invalid, or expired."),
    ("403", "API_KEY_REVOKED / API_KEY_EXPIRED", "Key exists but should not be used."),
    ("403", "IP_BLOCKED / THREAT_DETECTED", "Request is denied by security policy."),
    ("429", "RATE_LIMIT_EXCEEDED", "Redis policy says the caller has exceeded the limit."),
    ("502/503", "CONTROL_PLANE_UNAVAILABLE / POLICY_UNAVAILABLE / RATE_LIMIT_UNAVAILABLE", "Required security infrastructure is unavailable."),
], [0.75, 2.4, 4.05])

h(doc, "12. How Request Validation Works in Code")
p(doc, "Validation is both structural and semantic. Structural validation means JSON can be decoded into the expected Go structs. Semantic validation means the decoded values make sense for security decisions.")
add_table(doc, ["Layer", "Example", "Failure Result"], [
    ("Config", "BACKEND_URL must exist, use http/https, and include host.", "Process fails during startup."),
    ("HTTP header", "X-API-Key must be present after TrimSpace.", "401 API_KEY_MISSING."),
    ("Control Plane response", "data must not be null.", "502 CONTROL_PLANE_UNAVAILABLE."),
    ("API key status", "active allowed, revoked/expired denied.", "403 or 401 depending on state."),
    ("JWT", "Algorithm, signature, exp, issuer, audience.", "401 JWT_* or 502 policy unavailable."),
    ("Policy", "Regex compiles, limits are positive, algorithm supported.", "503/502 policy unavailable."),
], [1.2, 3.0, 2.7])

h(doc, "13. Go Concepts This Project Teaches")
bullets(doc, [
    "Interfaces: Handler depends on KeyValidator, JWTValidator, RateLimiter, IPBlocker, ThreatDetector, and EventReporter interfaces. That makes tests easy because fake implementations can be injected.",
    "Pointers: *JWTPolicy means the policy may be absent. *time.Time means expires_at may be null. *int64 lets events omit optional IDs.",
    "Context: context.WithTimeout gives downstream validation calls a deadline tied to the incoming request.",
    "Struct tags: json:\"api_key\" tells encoding/json how to map Go fields to JSON fields.",
    "Error wrapping: errors.Is works because code wraps sentinel errors like ErrKeyExpired and ErrPolicyUnavailable.",
    "ReverseProxy: httputil.NewSingleHostReverseProxy handles the mechanics of forwarding to the backend.",
])

h(doc, "14. Security Design Choices")
bullets(doc, [
    "Raw API keys are not used as cache keys; the cached validator hashes them first.",
    "The proxy removes API key and Authorization headers before forwarding.",
    "Policy checks fail closed when policy, Redis, or validation dependencies are unavailable.",
    "Security events are intentionally small and sanitized.",
    "The policy stream can disconnect without immediately disabling protection because the last valid snapshot remains available.",
])

h(doc, "15. Config Cheat Sheet")
add_table(doc, ["Variable", "Required", "Meaning / Default"], [
    ("BACKEND_URL", "yes", "Upstream service the proxy forwards to."),
    ("CONTROL_PLANE_URL", "yes", "Base URL for API key validation, policy fetch, and event reporting."),
    ("INTERNAL_API_TOKEN", "yes", "Shared secret sent as X-Aegis-Internal-Token."),
    ("LISTEN_ADDR", "no", "Default :8080."),
    ("API_KEY_HEADER", "no", "Default X-API-Key."),
    ("VALIDATION_TIMEOUT", "no", "Default 2s."),
    ("VALIDATION_CACHE_TTL", "no", "Default 30s for successful API-key validation cache."),
    ("VALIDATION_NEGATIVE_CACHE_TTL", "no", "Default 5s for negative validation cache."),
    ("POLICY_CACHE_TTL", "no", "Default 30s REST fallback policy cache."),
    ("REDIS_ADDR / REDIS_PASSWORD / REDIS_DB", "mixed", "Redis settings; addr defaults localhost:6379, DB defaults 0."),
    ("GRPC_POLICY_ADDR", "no", "Default localhost:50051."),
    ("SHUTDOWN_TIMEOUT", "no", "Default 10s."),
], [1.9, 0.8, 4.4])

h(doc, "16. Reading Path for Reverse Learning")
numbered(doc, [
    "Start with README.md to understand the product contract.",
    "Read internal/proxy/policy.go to understand the data structures.",
    "Read internal/proxy/handler.go and trace the ServeHTTP method line by line.",
    "Read internal/controlplane/client.go to see how JSON is encoded and decoded.",
    "Read jwt_validator.go, ip_blocker.go, threat_detector.go, and rate_limiter.go as independent security modules.",
    "Read cmd/proxy/main.go last to understand dependency wiring and startup/shutdown.",
    "Read tests after each file. Tests are the quickest way to see what behavior the author wanted.",
])

h(doc, "17. Practice Exercises")
bullets(doc, [
    "Change API_KEY_HEADER in the environment and observe which header the handler checks.",
    "Create a fake PolicySnapshot in a test and add a blocked IP address.",
    "Add a threat rule that matches GET /admin and watch the handler return THREAT_DETECTED.",
    "Compare fixed_window and token_bucket rate limiting by reading the Redis scripts.",
    "Add a new error case and trace how it becomes JSON through writeError.",
])

h(doc, "18. Source Map")
add_table(doc, ["Question", "Best File"], [
    ("Where is a request accepted or rejected?", "internal/proxy/handler.go"),
    ("What JSON fields does policy use?", "internal/proxy/policy.go"),
    ("How is an API key sent to the Control Plane?", "internal/controlplane/client.go"),
    ("How does JWT validation work?", "internal/proxy/jwt_validator.go"),
    ("How are rate limits selected and enforced?", "internal/proxy/rate_limiter.go"),
    ("How does app startup wire everything together?", "cmd/proxy/main.go"),
    ("Where are environment variables parsed?", "internal/config/config.go"),
], [3.0, 4.0])

h(doc, "Appendix: One-Page Mental Model")
p(doc, "Think of the project as five gates before forwarding:")
numbered(doc, [
    "Network gate: Is this IP blocked?",
    "Threat gate: Does method plus URI match a dangerous pattern?",
    "API key gate: Does the Control Plane recognize this key and say it is active?",
    "JWT gate: Is this bearer token signed correctly and not expired?",
    "Rate gate: Has this key, route, or global traffic exceeded its Redis-backed limit?",
])
p(doc, "Only after those gates pass does the proxy become boring in the best possible way: it forwards the request and lets the backend do backend work.")

doc.add_page_break()
h(doc, "Deep Dive A: Reverse Proxy Fundamentals")
p(doc, "A normal backend server receives a request and produces a response. A reverse proxy receives the request first, optionally changes or validates it, then sends it to the backend. The client usually does not know or care that another server is behind the proxy.")
p(doc, "In this project, httputil.NewSingleHostReverseProxy is the Go standard-library helper that does the forwarding. It rewrites the outbound request so it targets BACKEND_URL, sends the request to that upstream, then streams the upstream response back to the client.")
bullets(doc, [
    "Forwarding is the easy part. The important project logic happens before forwarding.",
    "The handler protects the backend from invalid credentials, blocked IPs, suspicious URLs, and too much traffic.",
    "The proxy strips X-API-Key and Authorization before the backend sees the request. That is a deliberate trust-boundary decision.",
    "If forwarding fails, the reverse proxy ErrorHandler returns BACKEND_UNAVAILABLE.",
])
p(doc, "Reverse learning tip: read ServeHTTP once for the happy path, then read it again while asking, 'What could go wrong at this line?' The code is mostly a chain of early returns.")

doc.add_page_break()
h(doc, "Deep Dive B: The Handler as a Security Pipeline")
p(doc, "Handler.ServeHTTP is the heart of the data plane. It is written as a pipeline: each guard either allows the request to continue or writes an error response and returns.")
add_table(doc, ["Order", "Guard", "Why It Runs Here"], [
    ("1", "IP block", "Cheap check. If the client IP is blocked, do not spend work parsing credentials."),
    ("2", "Threat detection", "Uses request method and URI. Blocks known suspicious request shapes early."),
    ("3", "API key presence", "A missing key is simple to reject before calling the Control Plane."),
    ("4", "API key validation", "Confirms the key exists, is active, and is not expired."),
    ("5", "JWT bearer parsing", "Ensures the Authorization header has exactly Bearer <token>."),
    ("6", "JWT validation", "Verifies algorithm, signature, exp, issuer, and audience from policy."),
    ("7", "Rate limiting", "Uses confirmed key identity and request path to choose a Redis bucket."),
    ("8", "Forwarding", "Only reached after every security decision passes."),
], [0.65, 1.8, 4.5])
p(doc, "This style is idiomatic Go server code: handle an error immediately, return immediately, and keep the main path easy to see.")

doc.add_page_break()
h(doc, "Deep Dive C: JSON Tags and Exported Fields")
p(doc, "Go only serializes exported struct fields, which means fields that start with a capital letter. JSON APIs often use snake_case. Struct tags bridge those two worlds.")
add_code(doc, 'type validationRequest struct {\n    APIKey string `json:"api_key"`\n}')
p(doc, "APIKey is exported, so encoding/json can see it. The json tag says the wire field should be api_key. Without the tag, Go would default to APIKey in JSON, which would not match the Control Plane contract.")
add_table(doc, ["Go Field", "JSON Field", "Why"], [
    ("GeneratedAt", "generated_at", "Public API style uses snake_case."),
    ("RateLimitRules", "rate_limit_rules", "A list of rules, not one rule."),
    ("BlockedIPAddresses", "blocked_ip_addresses", "Explicitly names exact blocked IPs."),
    ("ExpiresAt", "expires_at", "Can be null, so the Go type is *time.Time."),
    ("ErrorCode", "errorCode", "Existing error contract uses camelCase for error code."),
], [1.6, 1.8, 3.6])
p(doc, "Reverse learning tip: when you see a struct tag, ask: 'Is this field part of an external contract?' If yes, changing it is a breaking API change.")

doc.add_page_break()
h(doc, "Deep Dive D: Pointers, Null, and Optional Data")
p(doc, "Several fields are pointers because the value may be absent. This matters a lot when translating JSON into Go.")
add_table(doc, ["Field", "Type", "Meaning of nil"], [
    ("PolicySnapshot.JWT", "*JWTPolicy", "No JWT policy exists. The validator fails closed with policy unavailable."),
    ("KeyInfo.ExpiresAt", "*time.Time", "The API key has no expiration date."),
    ("RateLimitPolicy.BurstAllowance", "*int64", "No extra token-bucket burst capacity."),
    ("SecurityEvent.APIKeyID", "*int64", "The request failed before a key ID was known."),
    ("SecurityEvent.RuleID", "*int64", "The request was not tied to a specific threat rule."),
], [1.9, 1.4, 3.7])
p(doc, "A beginner Go habit worth building: whenever you see a pointer, ask whether nil is a valid business state or just an implementation detail. In this project, nil usually means 'optional' or 'not known yet'.")

doc.add_page_break()
h(doc, "Deep Dive E: API Key Validation Contract")
p(doc, "The proxy calls the Control Plane with a private internal token and a JSON body containing the client API key. The internal token authenticates the proxy itself. The API key identifies the external caller.")
add_code(doc, 'POST /api/v1/api-keys/validate\nContent-Type: application/json\nX-Aegis-Internal-Token: <shared internal token>\n\n{"api_key":"ak_..."}')
p(doc, "There are two secrets in this part of the flow, but they mean different things. INTERNAL_API_TOKEN is service-to-service trust. X-API-Key is client-to-platform trust.")
add_table(doc, ["Response Case", "Code Behavior", "Client Result"], [
    ("200 with active data", "Returns KeyInfo.", "Request continues."),
    ("200 with revoked status", "Returns ErrKeyRevoked.", "403 API_KEY_REVOKED."),
    ("200 with expired expires_at", "Returns ErrKeyExpired.", "403 API_KEY_EXPIRED."),
    ("404 or API_KEY_NOT_FOUND", "Returns ErrKeyNotFound.", "401 API_KEY_INVALID."),
    ("Control Plane unauthorized", "Wrapped as validator unavailable.", "502 CONTROL_PLANE_UNAVAILABLE."),
    ("Bad/missing data", "Wrapped as validator unavailable.", "502 CONTROL_PLANE_UNAVAILABLE."),
], [1.9, 2.3, 2.7])

doc.add_page_break()
h(doc, "Deep Dive F: Shape Validation vs Semantic Validation")
p(doc, "Shape validation asks whether JSON can be decoded. Semantic validation asks whether decoded values are acceptable for security decisions.")
add_table(doc, ["Question", "Example", "Where It Happens"], [
    ("Can JSON decode?", "Does data look like KeyInfo?", "json.Unmarshal in the Control Plane client."),
    ("Was data present?", "envelope.Data == nil", "ValidateKey."),
    ("Is status allowed?", "active or revoked", "ValidateKey switch statement."),
    ("Is the key expired?", "ExpiresAt must be after time.Now().", "ValidateKey."),
    ("Is duration positive?", "VALIDATION_TIMEOUT > 0", "config.durationFromEnv."),
    ("Is path valid?", "CONTROL_PLANE_VALIDATE_PATH starts with /", "config.Load."),
    ("Is policy usable?", "Regex compiles and limits are positive.", "Threat detector and rate limiter."),
], [1.8, 2.6, 2.6])
p(doc, "This is the biggest learning point in the project: decoding JSON into a struct does not mean the request is valid. It only means the request has the expected shape.")

doc.add_page_break()
h(doc, "Deep Dive G: JWT Validation")
p(doc, "JWT validation depends on policy because the proxy should not hard-code the signing algorithm, verification key, issuer, or audience. Those values can change centrally in the Control Plane.")
numbered(doc, [
    "The handler extracts Authorization.",
    "bearerToken requires two fields: Bearer and the token string.",
    "PolicyJWTValidator loads the latest policy snapshot.",
    "parseVerificationKey converts the configured key into the correct Go key type.",
    "jwt.Parse verifies the token using allowed methods and parser options.",
    "Expired tokens become ErrJWTExpired; other invalid tokens become ErrJWTInvalid.",
])
p(doc, "Algorithm enforcement is especially important. Without it, a system can accidentally accept a token signed with an unexpected method. This project checks the allowed methods and also compares token.Method.Alg() explicitly.")

doc.add_page_break()
h(doc, "Deep Dive H: Rate Limiting Algorithms")
p(doc, "Rate limiting answers: 'Is this caller allowed to make one more request right now?' The proxy uses Redis so multiple proxy instances can share the same limit state.")
add_table(doc, ["Algorithm", "How It Thinks", "Tradeoff"], [
    ("fixed_window", "Count requests in a fixed time bucket.", "Simple and fast, but can allow bursts around window boundaries."),
    ("sliding_window", "Keep recent timestamps in a sorted set.", "Smoother behavior, more Redis work."),
    ("token_bucket", "Refill tokens over time and spend one per request.", "Good burst handling, slightly more math."),
], [1.5, 3.0, 2.5])
p(doc, "Rule selection is intentionally deterministic: exact API key rule wins, then exact route rule, then global rule. That makes policy behavior explainable when multiple rules exist.")
add_code(doc, 'api_key rule -> route rule -> global rule -> no rule means allowed')

doc.add_page_break()
h(doc, "Deep Dive I: IP Blocking and Threat Detection")
p(doc, "IP blocking is exact-match based. It parses r.RemoteAddr, removes a port if present, normalizes brackets and IPv4-mapped addresses, then compares against blocked_ip_addresses.")
p(doc, "Threat detection is regex based. It combines the method and URL RequestURI into a single string, such as GET /v1/users?id=1. That lets policy authors match method, path, and query patterns.")
add_table(doc, ["Risk", "Current Behavior", "Learning Note"], [
    ("Invalid client address", "Returns CLIENT_IP_INVALID.", "The proxy fails the request instead of guessing."),
    ("Invalid blocked IP in policy", "Returns policy unavailable.", "A broken policy is treated as unsafe."),
    ("Invalid regex in policy", "Returns policy unavailable.", "Security rules must compile before use."),
    ("Threat match", "Returns THREAT_DETECTED and records rule_id.", "Events can later explain which rule fired."),
], [1.6, 2.4, 3.0])

doc.add_page_break()
h(doc, "Deep Dive J: Security Events")
p(doc, "Every request outcome becomes a small SecurityEvent. The event is useful for analytics and incident review, but it avoids sensitive data.")
add_code(doc, 'type SecurityEvent struct {\n    EventType  string `json:"event_type"`\n    SourceIP   string `json:"source_ip"`\n    APIKeyID   *int64 `json:"api_key_id"`\n    RuleID     *int64 `json:"rule_id"`\n    Method     string `json:"method"`\n    Path       string `json:"path"`\n    StatusCode int    `json:"status_code"`\n}')
bullets(doc, [
    "It includes path, not query string.",
    "It includes API key ID, not the raw API key.",
    "It includes rule ID only when a rule was involved.",
    "It does not include request body or arbitrary headers.",
    "The reporter uses a bounded queue so reporting does not block request handling indefinitely.",
])

doc.add_page_break()
h(doc, "Deep Dive K: Caching and Policy Snapshots")
p(doc, "The proxy uses caching in two important places: API key validation and policy snapshots. Caching reduces load on the Control Plane and lets the proxy keep enforcing the last known good policy during short network interruptions.")
p(doc, "The important security detail is what gets cached. The validator caches by SHA-256 digest of the raw key, not the raw key itself. The policy cache stores structured policy snapshots.")
add_table(doc, ["Cache", "Positive Case", "Negative Case"], [
    ("API key validation", "Active KeyInfo cached for VALIDATION_CACHE_TTL.", "Unknown/revoked/expired style failures can be cached for a shorter negative TTL."),
    ("REST policy fallback", "Last fetched policy reused for POLICY_CACHE_TTL.", "If unavailable and no valid cache exists, policy checks fail closed."),
    ("Streaming policy", "Latest valid gRPC update replaces current snapshot.", "If stream disconnects, last valid snapshot remains in memory while reconnecting."),
], [1.7, 2.7, 2.7])

doc.add_page_break()
h(doc, "Deep Dive L: Startup Wiring in main.go")
p(doc, "main.go is a dependency assembly file. It does not contain most business logic; it chooses concrete implementations and wires them into the Handler interfaces.")
numbered(doc, [
    "Load environment configuration.",
    "Create shared HTTP client for Control Plane calls.",
    "Create API key validation client and cached validator.",
    "Create REST policy client and cached policy provider.",
    "Create streaming policy provider and start gRPC subscriber.",
    "Create JWT validator, Redis rate limiter, IP blocker, threat detector, and event reporter.",
    "Create proxy handler and wrap it with recovery/request logging middleware.",
    "Start the HTTP server and wait for shutdown signal.",
])
p(doc, "This pattern is common in Go: constructors validate dependencies, return errors early, and main exits if the application cannot be safely started.")

doc.add_page_break()
h(doc, "Deep Dive M: Middleware and Recovery")
p(doc, "Middleware wraps the handler with behavior that should apply to every request. In this project, RequestLogger records request-level information and Recover catches panics so a single bug does not crash the whole process.")
p(doc, "The middleware chain is built around normal net/http interfaces. That is a nice beginner Go lesson: powerful web servers can be composed with one interface, http.Handler.")
add_code(doc, 'server := &http.Server{\n    Addr: cfg.ListenAddr,\n    Handler: middleware.Chain(proxyHandler, middleware.Recover(logger), middleware.RequestLogger(logger)),\n}')
p(doc, "Reverse learning tip: middleware is easier to understand if you draw it as layers around ServeHTTP. A request enters the outer layer first and reaches the proxy handler last.")

doc.add_page_break()
h(doc, "Deep Dive N: Testing Strategy")
p(doc, "The tests in this project are valuable because most behavior is interface-driven. The handler can be tested with fake validators, fake rate limiters, fake policy providers, and httptest servers.")
add_table(doc, ["Test Type", "What It Proves", "Files To Read"], [
    ("Handler tests", "Correct HTTP statuses, headers, forwarding, and credential stripping.", "internal/proxy/handler_test.go"),
    ("JWT tests", "Algorithm, expiration, issuer, audience, and key parsing behavior.", "internal/proxy/jwt_validator_test.go"),
    ("Rate limiter tests", "Rule precedence and Redis algorithm behavior.", "internal/proxy/rate_limiter_test.go"),
    ("Policy cache tests", "Snapshot copy behavior and fallback behavior.", "internal/controlplane/policy_cache_test.go"),
    ("Config tests", "Required env variables, URL parsing, durations, and defaults.", "internal/config/config_test.go"),
], [1.7, 3.0, 2.3])
p(doc, "A good reverse-learning exercise is to pick one errorCode, search for it in tests, then trace backward to the code branch that produces it.")

h(doc, "Deep Dive O: Common Beginner Questions")
add_table(doc, ["Question", "Answer"], [
    ("Why so many interfaces?", "They let the handler depend on behavior, not concrete implementations. That makes tests and future replacement easier."),
    ("Why fail closed?", "If security policy cannot be loaded or enforced, allowing traffic would bypass protection."),
    ("Why remove credentials before forwarding?", "The backend should not accidentally log, store, or rely on edge credentials."),
    ("Why use Redis for rate limits?", "In-memory counters would break when more than one proxy instance is running."),
    ("Why validate config at startup?", "A misconfigured security proxy should not start and accept traffic."),
    ("Why use context timeouts?", "External calls should not hang the request forever."),
], [2.3, 4.7])

h(doc, "Deep Dive P: Field-by-Field Payload Atlas")
add_table(doc, ["Payload", "Field", "Beginner Explanation"], [
    ("validationRequest", "api_key", "The secret presented by the client. Sent only to the private Control Plane route."),
    ("apiResponse", "status", "Envelope-level status, not the key lifecycle status."),
    ("apiResponse", "data", "The actual KeyInfo object. Must exist for success."),
    ("KeyInfo", "status", "Lifecycle state such as active or revoked."),
    ("JWTPolicy", "verification_key", "The key material used to check JWT signatures."),
    ("RateLimitPolicy", "scope_type", "Where the rule applies: one key, one route, or everything."),
    ("ThreatPolicy", "pattern", "RE2-compatible regex tested against method plus URI."),
    ("SecurityEvent", "event_type", "Machine-readable outcome such as request_forwarded or api_key_missing."),
], [1.6, 1.7, 3.7])

h(doc, "Deep Dive Q: How To Read ServeHTTP Without Getting Lost")
numbered(doc, [
    "Ignore the event object on your first pass and read only the if statements.",
    "Write down every return statement and the error code returned before it.",
    "Read the happy path from key extraction to forwarder.ServeHTTP.",
    "Read the deferred event reporting after you understand the branches.",
    "Trace one request by hand: blocked IP, missing API key, invalid JWT, rate limited, and successful request.",
])
p(doc, "Once you see that ServeHTTP is mostly guard clauses, the code becomes much less intimidating. It is a checklist implemented as Go code.")

h(doc, "Deep Dive R: Suggested Refactor Ideas for Learning")
p(doc, "These are not required fixes. They are learning exercises that would help you understand the project more deeply.")
bullets(doc, [
    "Add scope enforcement so KeyInfo.Scopes controls which routes a key can call.",
    "Add CIDR support to IP blocking and compare netip.Prefix with exact netip.Addr matching.",
    "Precompile threat regex rules when the policy snapshot is loaded instead of compiling on every request.",
    "Add a correlation/request ID to SecurityEvent while still avoiding sensitive headers.",
    "Add a small architecture diagram to README.md showing Control Plane, Redis, proxy, and backend.",
])
p(doc, "The best exercise is scope enforcement because it connects JSON payload design, request validation, tests, and error responses in one feature.")

doc.save(OUT)
print(OUT)
